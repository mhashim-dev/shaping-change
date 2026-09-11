#!/usr/bin/env python3
"""Build the team training / handover guide as a Word doc, written for non-coders
(with a "who does what" split so the one HTML/CSS person can be the technical lead).
    python3 scripts/build-training-doc.py  -> docs/Shaping-Change-Team-Training-Guide.docx
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GOLD = RGBColor(0xB0, 0x86, 0x1a); DARK = RGBColor(0x2A, 0x1F, 0x10); GREY = RGBColor(0x55, 0x4d, 0x40)
GREEN = RGBColor(0x1e, 0x7a, 0x3c); BLUE = RGBColor(0x1c, 0x5a, 0x9c); RED = RGBColor(0xb0, 0x1a, 0x1a)

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)


def shade(paragraph, fill='F3F1EB'):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    pPr.append(shd)


def H(text, size, color=DARK, before=14, after=4, bold=True):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size); r.font.color.rgb = color
    return p


def para(text, after=6, italic=False, color=None, size=10.5):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text); r.italic = italic; r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    return p


def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(3)
    if bold_lead:
        r = p.add_run(bold_lead + ' '); r.bold = True
    p.add_run(text)
    return p


def role(who):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3); p.paragraph_format.space_before = Pt(2)
    if who == 'any':
        r = p.add_run('WHO: Anyone on the team can do this.'); r.font.color.rgb = GREEN
    else:
        r = p.add_run('WHO: Best for your HTML/CSS person.'); r.font.color.rgb = BLUE
    r.bold = True; r.font.size = Pt(9.5)
    return p


def code(text):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(8); p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.left_indent = Pt(6)
    lines = text.split('\n')
    for i, ln in enumerate(lines):
        if i: p.add_run().add_break()
        r = p.add_run(ln if ln else ' '); r.font.name = 'Consolas'; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    shade(p)
    return p


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Light Grid Accent 1'
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]; c.text = ''
        rn = c.paragraphs[0].add_run(h); rn.bold = True; rn.font.size = Pt(9.5)
    for row in rows:
        cells = t.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = ''
            rn = cells[j].paragraphs[0].add_run(val); rn.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


# ---------------- COVER ----------------
title = doc.add_paragraph(); r = title.add_run('Shaping Change'); r.bold = True; r.font.size = Pt(26); r.font.color.rgb = DARK
sub = doc.add_paragraph(); r = sub.add_run('Team Training & Handover Guide'); r.bold = True; r.italic = True; r.font.size = Pt(16); r.font.color.rgb = GOLD
sub.paragraph_format.space_after = Pt(6)
para('How the game is built, and how the team can update it going forward — written for people who have not '
     'worked with web code before. No coding background is needed for the everyday work.', italic=True, color=GREY)

# who does what
H('Who does what on the team', 15, GOLD)
para('This game is set up so most updates need no coding at all. Here is a simple split of roles:')
bullet('edit the words on any screen, change or add an answer option, change the support links, and add new '
       'screens. This is careful list-editing in one file — not programming.', bold_lead='Anyone can:')
bullet('be the team’s “technical lead”: install the tools, run the short commands, rebuild the shareable file, '
       'and handle the “look” (colours, fonts, spacing — which is CSS). They are also the person to call if '
       'something breaks.', bold_lead='Your HTML/CSS person can:')
para('Throughout this guide, each task is tagged so you know who it’s for.', italic=True, color=GREY)

# 1
H('1.  The big picture (read this first)', 15, GOLD)
para('Shaping Change is a short web activity: the player follows Orion’s story and makes choices, and a tree '
     'grows or declines to reflect it. There are TWO copies of the same game — this matters:')
table(['Copy', 'What it’s for'],
      [['The dev app', 'What you edit and preview live while working (runs on your computer only).'],
       ['The offline game', 'The single file you SHARE — “Power-Pressure-Choice-POC.html”. It runs on any device with no internet.']])
para('Both copies read the SAME words and choices from one file, so you only write content once. The dev app '
     'updates instantly as you type. The offline game is “rebuilt” with one command when you’re ready to share.')
para('The most important idea: the WORDS AND CHOICES are kept in a separate file from the DRAWING AND '
     'MECHANICS. The team’s job is the words-and-choices file. The drawing engine can stay untouched.')

# 2
H('2.  A few words you’ll hear (in plain terms)', 15, GOLD)
bullet('the engine the app is built with. Think of it as the machinery that turns our files into a web page. '
       'You rarely touch this.', bold_lead='Next.js / React —')
bullet('the programming language. Our content file uses a very simple, list-like slice of it. That’s the part you edit.',
       bold_lead='JavaScript —')
bullet('a free program your computer needs so the commands work. Installed once.', bold_lead='Node.js / npm —')
bullet('the text window where you type commands. In the editor: View → Terminal.', bold_lead='Terminal —')
bullet('your own computer. When the app runs you preview it at http://localhost:3000 — private to you.', bold_lead='localhost —')
bullet('the free code editor we use to open the project and edit files (code.visualstudio.com).', bold_lead='VS Code —')

# 3
H('3.  One-time setup', 15, GOLD)
role('css')
para('Do these once on each computer that will edit the game:')
bullet('Install Node.js — go to nodejs.org, download the “LTS” version, install with the defaults.')
bullet('Install VS Code — from code.visualstudio.com.')
bullet('In VS Code: File → Open Folder → choose the “ames-tree-main” folder.')
bullet('Open the Terminal (View → Terminal) and run:')
code('npm install')
para('That downloads the project’s parts (a minute or two). Only needed once per computer.', italic=True, color=GREY)

# 4
H('4.  See the game and your changes', 15, GOLD)
role('any')
para('In the Terminal, run this and leave it running:')
code('npm run dev')
para('Then open a browser to http://localhost:3000. As you edit and save files, the page updates automatically. '
     'To stop it, click the Terminal and press Ctrl + C.')
para('To open the OFFLINE shareable game, just double-click “Power-Pressure-Choice-POC.html” — it opens in any '
     'browser. (It only shows your latest changes after you rebuild it — see section 9.)')

# 5
H('5.  Which file does what', 15, GOLD)
para('You’ll spend almost all your time in the first row.')
table(['File', 'What it holds', 'Who / how often'],
      [['lib/game-content.js', 'Every screen: the words, the choices, the order.', 'Anyone — often (your main file)'],
       ['app/globals.css', 'Colours, fonts, spacing (the “look”).', 'HTML/CSS person — sometimes'],
       ['lib/i18n.js', 'Fixed button/label wording; future-language setup.', 'Occasionally'],
       ['lib/tree-engine.js', 'Draws the tree, Orion, thought bubbles, weather.', 'Advanced — leave alone'],
       ['lib/audio-engine.js', 'Background music and sound cues.', 'Advanced — rarely'],
       ['scripts/validate-content.mjs', 'The safety check for the pedagogy.', 'You RUN it; don’t edit it'],
       ['scripts/build-poc.cjs', 'Builds the offline shareable game.', 'You RUN it; don’t edit it']])

# 6
H('6.  How one screen works', 15, GOLD)
para('Open lib/game-content.js. It’s a list called STEPS. Each screen in the game is one block in that list. '
     'Here is a real one (“Money worries”), with notes after each line:')
code("{\n"
     "  id: 'pressure_money',        // the screen's internal name\n"
     "  think: 'pressures',          // which picture shows above Orion (optional)\n"
     "  type: 'story',               // the KIND of screen (see below)\n"
     "  phase: 'Pressure',           // which chapter (drives the progress dots)\n"
     "  tag: 'Phase 1 · The pressures build',   // small gold line at the top\n"
     "  prompt: 'Money worries',     // the big heading\n"
     "  hint: 'Back home, Orion ...',// the body text the player reads\n"
     "  btn: 'Continue',             // the button label\n"
     "  kind: 'neutral',             // healthy | harmful | neutral (the meaning)\n"
     "  fx: { health: -7 },          // effect on the hidden 'health' score\n"
     "  next: 'pressure_role'        // which screen comes next (by its id)\n"
     "}")
para('The five kinds of screen (the “type”):')
bullet('a welcome / character screen.', bold_lead='intro —')
bullet('narration with one Continue button (its “next” points to the following screen).', bold_lead='story —')
bullet('a choice screen — instead of “next” it has a list of “options” (buttons).', bold_lead='pick —')
bullet('the pledge screen (the player writes or picks their own step).', bold_lead='commit —')
bullet('the final support screen.', bold_lead='ending —')
para('A choice screen’s options look like this — each block is one button:')
code("options: [\n"
     "  {\n"
     "    label: 'Take tight control of the money',  // the button text\n"
     "    kind: 'harmful',       // healthy | harmful | neutral\n"
     "    next: 'impact',        // where this choice leads\n"
     "    fx: { health: -18 },   // its effect on the tree\n"
     "    info: 'Orion’s reasoning: ...'  // feedback shown after they pick\n"
     "  },\n"
     "]")
para('How the story connects, and how the ending is chosen:')
bullet('The ORDER is set by “next” (on story screens) and by each option’s “next” (on choice screens). Every '
       'screen points to the next one by its id.')
bullet('There is a hidden “health” score (it starts at 68). Each “fx: { health: … }” nudges it up or down — '
       'that’s what makes the tree greener or barer.')
bullet('At the end, the game reads the health score to pick the outcome: 70+ = healthy tree, 40–69 = healing '
       'tree, below 40 = tree under strain.')

# 7
H('7.  Recipes — the common changes', 15, GOLD)
para('Golden rule for every change: after editing, run the safety check (section 8), then look at it in the browser.',
     italic=True, color=RED)

H('Recipe A — Change the words on a screen', 12, DARK, before=8, after=2)
role('any')
bullet('Open lib/game-content.js. Press Ctrl/Cmd + F and search for a phrase you can see on screen (or the id).')
bullet('Edit the text inside the quotes for prompt (heading), hint (body) or btn (button). Save. Check the browser.')
bullet('Inside “hint” text, \\n means “start a new line”.')

H('Recipe B — Change or add an answer option', 12, DARK, before=8, after=2)
role('any')
bullet('Find the screen (a type: pick), e.g. id: attitude.')
bullet('To edit: change its label and info. To add: copy an existing { … } option, paste it into the list '
       '(mind the commas), and set label, kind, next, fx and info.')
bullet('Keep the rules: a harmful option must never RAISE health; a help-seeking option should heal.')

H('Recipe C — Change the support links (last screen)', 12, DARK, before=8, after=2)
role('any')
bullet('Find id: support. Edit the “steps” list — each has a label (visible text) and a url (the link).')

H('Recipe D — Change a picture (thought bubble) on a screen', 12, DARK, before=8, after=2)
role('any')
bullet('Add, remove or change the “think:” value. Options: money · pressures · family_role · family · '
       'family_sad · family_happy · talk_together · partner_sad · children_sad.')

H('Recipe E — Add a whole new screen', 12, DARK, before=8, after=2)
role('any')
bullet('Copy an existing screen block, give it a NEW unique id, set its words.')
bullet('Wire it in: point the previous screen’s “next” to your new id, and set your new screen’s “next” to '
       'where it goes afterwards. Run the safety check — it flags broken links.')

H('Recipe F — Change the look (colours, fonts, spacing)', 12, DARK, before=8, after=2)
role('css')
bullet('These live in app/globals.css — this is CSS, so it’s the natural home for your HTML/CSS person. '
       'Search for a colour or size and adjust. Safe to experiment — undo (Ctrl/Cmd + Z) if it looks wrong.')

# 8
H('8.  The safety check (very important)', 15, GOLD)
role('any')
para('Because this is a family-violence PREVENTION tool, some rules must never be broken by accident. After any '
     'content change, run:')
code('npm run test:content')
para('It confirms, among other things: every screen is reachable and its links are real; a harmful choice can '
     'NEVER heal the tree; help-seeking always heals; pressures/impacts are never a “wrong answer”; all three '
     'outcomes stay reachable; the strained outcome is always reversible; the 1800RESPECT / 000 support line is '
     'on every screen; and statistics keep their sources.')
para('If it prints “✓ All 10 invariant groups hold,” you’re good. If it prints an error, it names exactly what '
     'to fix. Always run this before sharing.')

# 9
H('9.  The everyday workflow', 15, GOLD)
code('1. Edit lib/game-content.js and Save\n'
     '2. npm run test:content   (safety check — fix anything it flags)\n'
     '3. Look at http://localhost:3000   (needs "npm run dev" running)\n'
     '4. npm run build:poc      (updates the offline shareable game)\n'
     '5. Share Power-Pressure-Choice-POC.html')
role('any')
para('Steps 1–3 are for anyone. Steps 4–5 (the rebuild) are a good fit for your HTML/CSS person, though the '
     'command is a single line anyone can run.', italic=True, color=GREY)

# 10
H('10.  Making the shareable files', 15, GOLD)
role('css')
bullet('The offline game (send this to anyone): run “npm run build:poc”, then share '
       'Power-Pressure-Choice-POC.html. One self-contained file, no internet, nothing is sent anywhere — '
       'answers stay on the player’s device.')
bullet('The review Word doc & deck (the screen-by-screen pack): these are rebuilt with separate scripts and '
       'need the screen images re-rendered first. This is the one task worth pairing with someone technical the '
       'first time — ask for a walkthrough.')

# 11
H('11.  Be careful with these', 15, GOLD)
bullet('Don’t rename an id unless you also update every “next” that points to it (the safety check catches this).')
bullet('Watch your commas and quotes in game-content.js. A missing comma or quote will make the page go blank — '
       'if that happens, undo your last edit (Ctrl/Cmd + Z).')
bullet('Never remove the support line, and never let a harmful choice heal the tree. The safety check guards '
       'this — respect it.', )
bullet('Any change to safety-message wording or statistics should go back to Edu / PVAW before publishing.')

# 12
H('12.  Troubleshooting', 15, GOLD)
table(['Symptom', 'Likely fix'],
      [['The page is blank', 'A typo in game-content.js (missing comma/quote). Undo your last change; re-save.'],
       ['Changes don’t show in the browser', 'Is “npm run dev” running? Refresh the page.'],
       ['Changes don’t show in the offline file', 'Run “npm run build:poc” — the offline copy only rebuilds on command.'],
       ['“npm: command not found”', 'Node.js isn’t installed (section 3).'],
       ['The old screen keeps showing', 'The game remembers progress. Click “Start over”.'],
       ['The safety check fails', 'Read the message — it names the exact rule and screen to fix.']])

# 13
H('13.  Command cheat-sheet', 15, GOLD)
code('npm install            one-time setup on a new computer\n'
     'npm run dev            start the live preview (http://localhost:3000)\n'
     'npm run test:content   run the pedagogy safety check (after every edit)\n'
     'npm run build:poc      rebuild the offline shareable game')
para('The two files to remember: lib/game-content.js (all the words and choices) and '
     'scripts/validate-content.mjs (the rules). Start there.', italic=True, color=GREY)

out = os.path.join(ROOT, 'docs', 'Shaping-Change-Team-Training-Guide.docx')
doc.save(out)
print('saved', os.path.relpath(out, ROOT))
