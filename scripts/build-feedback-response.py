#!/usr/bin/env python3
"""Point-by-point response to EVERY note in the PAG feedback document (body notes + the 9
margin comments), for sharing back to the advisory group. Plain checklist — one line per note,
with a status and what was done.
    python3 scripts/build-feedback-response.py  -> docs/Shaping-Change-PAG-Response.docx
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GOLD = RGBColor(0xB0, 0x86, 0x1a); DARK = RGBColor(0x2A, 0x1F, 0x10); GREY = RGBColor(0x55, 0x4d, 0x40)
DONE = RGBColor(0x1e, 0x7a, 0x3c); DESIGN = RGBColor(0x8a, 0x5a, 0x00)
FUTURE = RGBColor(0x60, 0x60, 0x60); STEER = RGBColor(0xb0, 0x1a, 0x1a); INFO = RGBColor(0x1c, 0x5a, 0x9c)
COL = {'Done': DONE, 'Designer': DESIGN, 'Future': FUTURE, 'Your call': STEER, 'Answered': INFO, 'Process': GREY}

# section -> [ (status, point, what we did) ]
DATA = [
 ('Sound & audio', [
   ('Done', 'Align audio with the visual outcomes; positive = uplifting, less-healthy = sombre.',
    'Gentle outcome cues: a soft rising sound for healthier choices, a soft falling sound for less-healthy ones.'),
   ('Done', 'Softer, more subtle background music.', 'Audio kept minimal and unobtrusive.'),
   ('Done', 'Option to turn sound on or off.', 'Already on every screen.'),
 ]),
 ('User reflection & participation', [
   ('Future', 'Let players enter their own challenges / lived experiences.', 'Planned as a future version (Ali: drop-down options).'),
   ('Future', 'End-of-game space to share experiences (racism, discrimination, settlement).', 'Planned for a future version (Ali: “to be added”).'),
   ('Done', 'Reflective questions + download responses as a PDF (viewed positively).', 'Kept — the four reflection questions and the printable take-home summary remain.'),
 ]),
 ('User experience & navigation', [
   ('Done', 'Make it clearer when the player is Orion vs. a community leader.', 'Kept the focus on Orion throughout; dropped the “community leader” framing (per Ali).'),
   ('Done', 'Strengthen the transition into the rebuild / reflection.', 'Rebuild reframed as “Help Orion rebuild.”'),
   ('Done', 'Clarify why the optional question appears on the welcome page.', 'The readiness question was removed (per Ali).'),
   ('Done', 'Explain that some options are realistic but unhealthy, not “correct” answers.', 'Added a line to the behaviour screen.'),
   ('Done', '“He digs in” may not be understood.', 'Reworded to “He holds on to the old ways.”'),
   ('Done', 'Reduce the text on the opening screen.', 'Welcome copy trimmed.'),
   ('Done', 'The black cloud visual could be more prominent.', 'The stress cloud above Orion is now bigger and darker, with faint rain when he is most stressed — still no lightning (PVAW).'),
   ('Done', 'Use the thought-bubble device more throughout.', 'Added a combined thought bubble on the “Two big pressures” and behaviour screens — the money worry and the family shown together.'),
   ('Future', 'Character customisation (hairstyle, clothing, colour).', 'Planned for a future version.'),
 ]),
 ('Slide 3 — Money worries', [
   ('Done', 'Orion appears passive — show him actively searching for work, repeatedly rejected because his overseas qualifications are not recognised (structural, not personal failure).',
    'Rewritten: Orion is an engineer applying for jobs; the doors keep closing (qualifications not recognised, employers want local experience, few contacts) — named as barriers common for new arrivals, not his fault.'),
   ('Done', 'Include references to employment barriers.', 'Three of Ali’s barriers named on screen; language and financial pressure appear on nearby screens.'),
   ('Done', 'Replace “head of the family” with “his role in the family changes.”', 'Changed on the “A changing role” screen; “head of the family” is gone everywhere.'),
 ]),
 ('Slide 4 — A changing role', [
   ('Done', 'Draw Orion in the image but smaller than his wife and children.', 'His thought bubble now shows his wife and child full-size and Orion smaller — his perceived shrinking role.'),
   ('Done', 'Add icons / visual supports for lower-English readers (migration, loss of status).', 'Carried by the concrete thought bubbles — the money worry (migration/work pressure) and the smaller-Orion bubble (loss of status). Standalone icons can still be added by a designer if wanted.'),
 ]),
 ('Slide 5 — Two pressures', [
   ('Done', 'Show the two pressures together, and add “Orion is not sure where to go and he has a choice to make.”',
    'Combined onto one screen, the line was added, and both pressures now show together in one thought bubble (the money worry + the family).'),
 ]),
 ('Narrative & context', [
   ('Done', 'Keep employment challenges visible throughout, not overshadowed by migration.', 'The money screen was strengthened and kept prominent.'),
   ('Done', 'Reduce text, particularly for mobile.', 'Copy trimmed; the layout already fits every screen size without scrolling.'),
   ('Your call', 'Add a help-seeking option earlier to normalise help-seeking.',
    'Help-seeking is normalised in the belief feedback and is a rebuild option. Adding a healthy option to the behaviour screen would change its “all unhealthy” design — we’d like your steer before doing that.'),
   ('Your call', 'Provide belief examples / context before Step 3.',
    'Beliefs are currently surfaced after the behaviour (on the belief screen). Moving them before would be a reorder — we’d like your steer.'),
 ]),
 ('Slide 6 — Behaviour', [
   ('Done', 'Include an “Orion’s reasoning” explanation for each response option.',
    'Each of the three options now shows Orion’s reasoning (why he might act this way) alongside the effect on the family, keeping “understanding is not excusing.”'),
   ('Done', '“Who feels it?” → “Family experiences.”', 'Retitled.'),
 ]),
 ('Slide 8 — The belief', [
   ('Done', 'Include additional belief options.', 'Added Ali’s beliefs (“…strong and in control,” “men shouldn’t ask for help — private problems should stay private”); five options now.'),
   ('Answered', 'Do the three response pathways lead to the same page or different outcomes?',
    'They converge — all three behaviour choices lead to the same next screen. The branching into different outcomes happens later, in the rebuild.'),
   ('Done', 'Strengthen the message that pressures stay the same but responses can differ (Slide 9).',
    'Added to the belief screen: “the pressures stay the same — but the belief behind our choices shapes how we respond.”'),
 ]),
 ('Slides 10–15', [
   ('Done', 'Slide 10 — keep the main character; remove external leadership.', 'Done (Orion focus).'),
   ('Done', 'Slide 11 — add “making decisions together while remaining responsible.”', 'Added to the rebuild action.'),
   ('Done', 'Slide 12 — show the benefits of healthier choices (confidence, less stress, stronger relationships).', 'Added to “the family feels the change.”'),
   ('Done', 'Slide 13 — make the key message explicit (pressures continue, but support builds healthier relationships); can this be shown via images?', 'The message is on the healthy outcome, and the rebuild-action screen now pictures the healthy way — two people working things out together over a shared plan.'),
   ('Done', 'Slide 15 — small changes still make a meaningful difference.', 'Added to the healing outcome.'),
 ]),
 ('Slides 18–22', [
   ('Done', 'Slide 18 — review the wording.', 'Reworded (“holds on to the old ways”).'),
   ('Done', 'Slide 20 — the reflection questions need changing.', 'The four questions were revised.'),
   ('Done', 'Slide 22 — change the name; make commitments specific and time-bound.',
    'Renamed “One step you will take”; pledge is time-bound (“This month, I will…”).'),
 ]),
 ('Whole document', [
   ('Done', 'Check grammar and structure; keep the focus on Orion.', 'Wording made consistent; Orion-only framing throughout.'),
   ('Process', 'Conduct user testing with men’s groups before launch.', 'For your team to arrange.'),
 ]),
]


doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)

p = doc.add_paragraph(); r = p.add_run('Shaping Change — Response to PAG Feedback (point by point)')
r.bold = True; r.font.size = Pt(17); r.font.color.rgb = DARK
sub = doc.add_paragraph(); sub.paragraph_format.space_after = Pt(6)
rs = sub.add_run('Every note in the feedback document — the body notes and the 9 margin comments — with its '
                 'status and what was done. Status: Done · Designer (needs illustration) · Future (future version) · '
                 'Your call (needs your steer) · Answered · Process.')
rs.italic = True; rs.font.size = Pt(9.5); rs.font.color.rgb = GREY

for section, items in DATA:
    h = doc.add_paragraph(); h.paragraph_format.space_before = Pt(11); h.paragraph_format.space_after = Pt(3)
    hr = h.add_run(section); hr.bold = True; hr.font.size = Pt(13); hr.font.color.rgb = GOLD
    for status, point, did in items:
        b = doc.add_paragraph(style='List Bullet'); b.paragraph_format.space_after = Pt(4)
        sr = b.add_run(status + ' — '); sr.bold = True; sr.font.color.rgb = COL.get(status, DARK)
        b.add_run(point + '  ')
        dr = b.add_run('→ ' + did); dr.italic = True; dr.font.color.rgb = GREY

n = sum(len(i) for _, i in DATA)
out = os.path.join(ROOT, 'docs', 'Shaping-Change-PAG-Response.docx')
doc.save(out)
print('saved', os.path.relpath(out, ROOT), '(%d points)' % n)
