#!/usr/bin/env python3
"""Build the Phase 1 demo crib sheet as a Word doc.
    python3 scripts/build-phase1-demo-reference.py  -> docs/Phase-1-Demo-Reference.docx
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GOLD = RGBColor(0xB0, 0x86, 0x1a); DARK = RGBColor(0x2A, 0x1F, 0x10); GREY = RGBColor(0x55, 0x4d, 0x40)

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)


def heading(text, size, color=DARK, bold=True, space_before=10, space_after=4, italic=False):
    p = doc.add_paragraph(); p.space_before = Pt(space_before); p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size); r.font.color.rgb = color
    return p


def rule():
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(6)
    r = p.add_run('_' * 66); r.font.color.rgb = RGBColor(0xCC, 0xC4, 0xB6)


def bullet(lead, rest=''):
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(3)
    if lead:
        r = p.add_run(lead); r.bold = True
    if rest:
        p.add_run((' ' if lead else '') + rest)


def screen(title, phase, see, point):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title); r.bold = True; r.font.size = Pt(11.5); r.font.color.rgb = DARK
    rp = p.add_run('   ·   ' + phase); rp.italic = True; rp.font.size = Pt(9.5); rp.font.color.rgb = GOLD
    ps = doc.add_paragraph(); ps.paragraph_format.space_after = Pt(1); ps.paragraph_format.left_indent = Pt(10)
    ls = ps.add_run('See:  '); ls.bold = True; ps.add_run(see)
    pt = doc.add_paragraph(); pt.paragraph_format.space_after = Pt(2); pt.paragraph_format.left_indent = Pt(10)
    lt = pt.add_run('Point:  '); lt.bold = True; pt.add_run(point)


# ---- header ----
heading('Shaping Change — Phase 1 Demo: Screen-by-Screen Reference', 17, DARK, space_before=0, space_after=2)
sub = doc.add_paragraph(); sub.paragraph_format.space_after = Pt(4)
rs = sub.add_run('A quick crib sheet for the advisory-group demo. Phase 1 (screens 1–10) is what you walk through; '
                 'Phase 2 (11–24) is summarised at the end as "what comes next."')
rs.italic = True; rs.font.size = Pt(9.5); rs.font.color.rgb = GREY
rule()

# ---- framing ----
heading('Say these if challenged (the framing)', 13, GOLD)
for lead, rest in [
    ('Understanding patterns — not blaming people or cultures.', 'Strengths-based and non-judgemental.'),
    ('No depicted violence.', 'Carried by feeling, a thought bubble, and the tree — never a graphic scene.'),
    ('The 1800RESPECT / 000 support line is on every single screen.', ''),
    ('No right or wrong answers, and nothing is "game over."', 'Every path, including the hard one, is reversible.'),
    ('Primary prevention — men as leaders of positive change,', 'not as a problem.'),
    ('Built closely with Education & PVAW.', ''),
    ('Made for this audience:', 'simple English (EAL), about 10–15 minutes, self-paced, works on any device, '
     'and private — answers stay on the person’s own device.'),
]:
    bullet(lead, rest)

# ---- phase 1 ----
rule()
heading('Phase 1 — Experience the problem  (the demo)', 13, GOLD)
P1 = [
 ('1. Welcome — what this is', 'Opening',
  'Title "Shaping Change: Building strong roots for safety," plus the full plain-English intro — what the '
  'activity is, that settling in a new country brings pressures we can lead through, that you follow one '
  'person’s story with a tree that grows alongside, and "no right or wrong answers, no blame." A short '
  '"What to expect" list and an optional readiness check-in. Button: Begin.',
  'Sets the warm, non-judgemental tone and expectations up front, so the next screen is purely about meeting '
  'the character.'),
 ('2. Meet Orion', 'Opening',
  'Orion arrived in Australia in 2023; back home he was a respected community leader people came to for help. '
  'Now he is starting again. He stands by a mature tree and waves hello.',
  'Introduces him warmly — not a villain, someone many in the room will recognise. The player is a supportive guide.'),
 ('3. Money worries', 'The pressures build',
  'No work since arriving, money is tight, the home feels too small. The tree shows early strain; a thought '
  'bubble shows coins with a falling arrow.',
  'The first external pressure — a circumstance, not a fault. Pressure is real and understandable, never a "wrong answer."'),
 ('4. A changing role', 'The pressures build',
  'His wife wants to study and work; the children want their own paths. He feels less respected and listened '
  'to — losing his place as head of the family. Bubble shows his partner and children.',
  'Names loss of role and status — a key driver for many newly-arrived men. Builds empathy without excusing behaviour.'),
 ('5. The first pressure — migration', 'The pressures build',
  'Migration pressure named plainly — new language, hard to find work, money stress, few connections. Kept '
  'to a couple of short sentences.',
  'Split into one-pressure-per-screen (team request) so the text is lighter for EAL learners. Makes migration '
  'pressure explicit and recognisable.'),
 ('6. The second pressure — loss of status', 'The pressures build',
  'Back home he was respected and led others; here he can feel unseen, like starting from zero — which brings '
  'frustration and shame. Closes with: "these pressures are not his fault, but they shape how he feels and acts."',
  'Names loss of status plainly and reinforces the "not his fault, but shapes his actions" message, without '
  'overwhelming the reader.'),
 ('7. The behaviour under pressure', 'CHOICE',
  '"The pressure is heavy. What does Orion do?" Three realistic responses — take tight control, go quiet and '
  'pull away, or let the stress come out as anger. No perfect answer; gentle feedback follows each.',
  'The player witnesses the problem behaviour by choosing how a stressed person might act. Naming it is the '
  'first step — understanding why is not the same as excusing it.'),
 ('8. Who feels it?', 'The impact',
  'His partner feels stressed and unheard; the children grow quiet and pull away; Orion feels more alone, not '
  'seeing the effect. Bubble shows the family unhappy; the tree is visibly declining.',
  'Makes the ripple on the whole family concrete and visible — without depicting any violence.'),
 ('9. The belief underneath', 'CHOICE',
  '"What belief might sit behind this?" The player picks the belief that fits — e.g. "As the father, it is my '
  'job to lead and decide." No wrong choice.',
  'Surfaces the belief driving the behaviour — the "soil." A reflection point, not blame, that sets up the '
  'idea that changing beliefs can change actions.'),
 ('10. The whole tree', 'The tree model — closes Phase 1',
  'The reveal — roots = life pressures, soil = beliefs, trunk = behaviour, branches = impact on family. The '
  'player taps each part of the tree to discover it (the legend fills in; also keyboard-selectable; a gentle '
  'one-time hint points to the tree).',
  'The centrepiece metaphor. To change the outcome you look at the whole tree — and values can be kept and '
  'expressed in a healthier way. This is the pivot into the rebuild.'),
]
for t, ph, see, point in P1:
    screen(t, ph, see, point)

# ---- phase 2 ----
rule()
heading('What comes next — Phase 2  (in development; same tree, learner as leader)', 13, GOLD)
for lead, rest in [
    ('11. Now you lead —', 'the player steps in as the leader. The pressures remain, but new beliefs and actions can grow a healthier tree.'),
    ('12. A healthier belief (choice — drag onto the soil) —', 'pick a healthier belief: respect earned through trust and listening; worth beyond status or work; children safe to share feelings; or staying in charge if fair.'),
    ('13. A leader’s action (choice — drag onto the trunk) —', 'decide together; ask for help (English classes, job support, counselling); or keep control but be kinder. The action decides which tree grows.'),
    ('14. The family feels the change —', 'sharing power and asking for help: the partner feels respected, the children open up, Orion feels less alone. Bubble now shows the family smiling.'),
    ('15–17. Outcomes —', 'healthy (full green tree, bearing fruit), healing (mostly green, some strain — "every step counts"), or under strain (old patterns held — framed as "a warning, not the end," always offering "Go back and rebuild").'),
    ('18–20. Optional escalation path —', 'if he digs in: raised voice / ends the conversation → slams the door / walks away → a calm, serious warning of where unaddressed patterns lead. Non-physical, non-graphic, reversible at every step — "it is never too late to choose a different way."'),
    ('21. What we learn —', 'belief → action → a safe, respectful home, with four reflection questions.'),
    ('22. Bigger than one family —', 'zooms out to the wider issue; the statistics sit behind an opt-in "Show the numbers" button, each with its source (AIHW 2018; PM&C 2023; ABS 2021). "Men can lead the change."'),
    ('23. Your leadership pledge —', 'the player writes one small, real leadership action of their own (or picks a suggestion).'),
    ('24. You can keep growing —', 'concrete next steps (AMES "Stop Violence Against Women" Leadership Course, Neighbourhood Houses, men’s groups, community/faith groups), the key messages, the support line, and the option to print or save a personalised summary of their journey. An optional before/after reflection shows how their self-rating moved.'),
]:
    bullet(lead, rest)

# ---- Q&A ----
rule()
heading('Likely questions — quick answers', 13, GOLD)
for lead, rest in [
    ('"Isn’t this blaming migrant men / cultures?"', 'No — the whole frame is understanding patterns, not blaming people or cultures. Orion is a respected leader under real pressure; we separate understanding a behaviour from excusing it.'),
    ('"Is any violence shown?"', 'No. The impact is carried by feeling, a thought bubble, and the declining tree. The optional harder path uses recognisable non-physical behaviours, kept calm and non-graphic.'),
    ('"What if someone taps the ‘wrong’ answer?"', 'There is no wrong answer and no "game over." Feedback is gentle, and the strained path always offers a way back to rebuild.'),
    ('"Where’s the evidence / who signed off?"', 'Developed with Education & PVAW; the closing statistics are cited and are being source-verified before public launch.'),
    ('"How long, and who’s it for?"', 'About 10–15 minutes, self-paced, simple English, for newly-arrived and migrant-community men — engaged as leaders of positive change.'),
]:
    bullet(lead, rest)

foot = doc.add_paragraph(); foot.paragraph_format.space_before = Pt(10)
rf = foot.add_run('Support is always available — 1800RESPECT (1800 737 732); in an emergency, 000.')
rf.italic = True; rf.font.size = Pt(9.5); rf.font.color.rgb = GREY

out = os.path.join(ROOT, 'docs', 'Phase-1-Demo-Reference.docx')
doc.save(out)
print('saved', os.path.relpath(out, ROOT))
