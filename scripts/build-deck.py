#!/usr/bin/env python3
"""Build a clickable PowerPoint walkthrough from the screen mockups.

A title slide, one full-bleed screen per stage (with speaker notes), and a closing
"key messages" slide. Click-to-advance is disabled, so navigation happens ONLY via
the on-slide buttons — which lets the journey branch at the Phase-2 rebuild fork and
loop back through the three outcomes. Run in Slide Show mode to click through it.

Mirrors the two-phase game (Edu team's final Blueprint — Orion):
  Phase 1  Experience the problem — the tree declines  (intro → pressures → behaviour
           → impact → belief → the whole-tree reveal)
  Phase 2  Rebuild the outcome — the tree heals or stays strained  (rebuild → attitude
           → behaviour → impact → one of three outcomes → takeaways → leadership → support)

    pip install python-pptx
    python3 scripts/build-deck.py        # -> docs/Power-Pressure-Choice-Deck.pptx
"""
import os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMG = os.path.join(ROOT, 'docs', 'figma-screens')
HERO = os.path.join(ROOT, 'docs', 'images', 'scene-flourishing.png')

SCREENS = [
    '01-intro.png', '02-pressure-money.png', '03-pressure-role.png',
    '04-pressures-explained.png', '05-behaviour.png', '06-impact.png', '07-belief.png',
    '08-root-tree-model.png', '09-rebuild-intro.png', '10-attitude.png',
    '11-behaviour-good.png', '12-impact-good.png', '13-outcome-healthy.png',
    '14-outcome-mixed.png', '15-outcome-damaged.png', '16-insist-1.png',
    '17-insist-2.png', '18-escalation-warning.png', '19-takeaways.png',
    '20-why-matters.png', '21-leadership.png', '22-support.png',
]

# slide indices: 0 title · 1..22 screens · 23 closing
TITLE, FIRST, CLOSING = 0, 1, 23

# slide index -> list of (label, target index, style)   style: go | harm | sub
NAV = {
    0:  [('Begin  ▶', 1, 'go')],
    1:  [('Continue  ▶', 2, 'go')],
    2:  [('Continue  ▶', 3, 'go')],
    3:  [('Continue  ▶', 4, 'go')],
    4:  [('Continue  ▶', 5, 'go')],
    5:  [('See the impact  ▶', 6, 'go')],
    6:  [('Continue  ▶', 7, 'go')],
    7:  [('Show the whole tree  ▶', 8, 'go')],
    8:  [('Now, help him rebuild  ▶', 9, 'go')],
    9:  [('Continue  ▶', 10, 'go')],
    10: [('Continue  ▶', 11, 'go')],
    # Phase-2 rebuild fork — how Orion now acts decides which tree grows.
    11: [('Share power & seek support  ▶', 12, 'go'),
         ('Change only a few things  ▶', 14, 'sub'),
         ('Keep the old pattern  ▶', 15, 'harm')],
    12: [('See the outcome  ▶', 13, 'go')],
    13: [('Key messages  ▶', 19, 'go')],
    14: [('←  Go back & rebuild', 9, 'sub'), ('Key messages  ▶', 19, 'go')],
    # Under strain → the optional escalation branch (he digs in) or go back and rebuild.
    15: [('←  Go back & rebuild', 9, 'go'), ('But he digs in  ▶', 16, 'harm')],
    16: [('←  Go back & rebuild', 9, 'go'), ('He refuses to listen  ▶', 17, 'harm')],
    17: [('←  Go back & rebuild', 9, 'go'), ('He will not back down  ▶', 18, 'harm')],
    18: [('←  Go back & rebuild', 9, 'go'), ('What this teaches us  ▶', 19, 'sub')],
    19: [('Why this matters  ▶', 20, 'go')],
    20: [('Your turn to lead  ▶', 21, 'go')],
    21: [('Continue  ▶', 22, 'go')],
    22: [('Key messages  ▶', 23, 'go'), ('↺  Play again', 1, 'sub')],
    23: [('↺  Start over', 1, 'sub'), ('Replay the rebuild', 9, 'sub')],
}

NOTES = {
    0:  'Open with the core idea: the same pressures can grow very different trees — what changes the outcome is belief, response and leadership. This is a two-part story: first we watch the problem grow, then we help rebuild. Click Begin to start.',
    1:  'Meet Orion. He arrived in Australia in 2023. Back home he was a respected community leader; now he is starting again — new country, new language, new life. We follow his story, then help him make better choices. Frame it as understanding patterns, not blaming people or cultures.',
    2:  'The first pressure: money. He has not found work since arriving, money is tight, and he thinks his home is too small for the family. A real, external pressure — not his fault and not a wrong answer.',
    3:  'The second pressure: loss of role and status. People used to come to him; now he feels unseen. Naming the pressure is not excusing what comes next.',
    4:  'Dedicated explainer (added at the Edu team\'s request): names the two big pressures — migration pressure and loss of status — tied to Orion, and stresses they are not his fault. Kept as an interactive screen (videos are out of scope for this release, due to timeline).',
    5:  'Phase 1, the behaviour. Under pressure he takes tight control — of the money, the decisions, the household. We are witnessing the problem here, on purpose, so we can understand it.',
    6:  'The impact reaches the whole family: his partner feels unheard, the children grow quiet and pull away, and he does not see how his behaviour is affecting the family he cares about. The tree is visibly declining now.',
    7:  'The belief underneath: "As the father, it is my job to lead and decide." A reflection point — where does that belief come from, and what does it cost?',
    8:  'The reveal: the whole tree. Roots are life pressures, soil is beliefs, trunk is behaviour, branches are the impact on family. Our values and beliefs shape how we act — so to change things we look at the whole tree, and we can keep our values while showing them in a healthier way. This closes Phase 1.',
    9:  'Phase 2 begins: now we help Orion rebuild. The pressures have not disappeared — but the response can change.',
    10: 'A new attitude is the first rebuild choice — the soil. A healthier belief (for example, "respect is earned through trust, not control," or wanting his children to feel safe sharing their feelings) starts to heal the tree.',
    11: 'This is the hinge of Phase 2. How he acts now decides which tree grows. Share power and seek support → it heals. Change only a little → it half-heals. Keep the old pattern → it stays under strain. Click one to see that outcome — you can always come back and try another.',
    12: 'The healthy behaviour lands: talking about finances and big decisions together, help sought, power shared. The home feels calmer and safer.',
    13: 'A healthy tree — earned, not by luck — now bearing fruit. A different belief, shared power, and asking for support. The pressures did not vanish; he learned to face them with help.',
    14: 'A tree that is healing — some things got better, some strain remains. Change is a journey; every step counts. He can go back and rebuild further.',
    15: 'A tree under strain — the old patterns held on. This is a warning, never "game over": he can go back and rebuild at any time. Or, to show where this can lead, follow the optional escalation path (he digs in).',
    16: 'Optional escalation (reworked with PVAW): he digs in — raises his voice and ends the conversation. The family goes quiet and pulls further away. This is his behaviour escalating, shown without any external "spark," and it is always reversible.',
    17: 'The escalation continues — he slams the door and walks away, arms crossed. "I\'m only trying to protect them," but control is not protection. Kept calm and non-graphic; the way back to rebuild is always offered.',
    18: 'The warning beat: "Where this can lead." Left unaddressed, these patterns harden into less safety at home — how everyday stress, unmanaged, can grow into harm. Deliberately calm and serious (never frightening), with no fire or lightning: the cause is internal, and it is never too late to choose a different way.',
    19: 'Land the takeaways: the tree can always recover, and reflect back through the tree — belief, leader thinking, behaviour, impact.',
    20: 'Why this matters — zoom out from one family to the bigger picture. Per Anu, the statistics sit here at the END (not the start), behind a "Show the numbers" button: now that players understand the causes, the numbers land as "why this matters" rather than a scare. Figures are finalised with citations (AIHW 2018 · PM&C 2023 · ABS 2021). Close on the hopeful note — men can lead the change.',
    21: 'The player writes one real, small leadership action of their own — that is what makes it theirs.',
    22: 'Support and next steps: the AMES SVAW course and community services, with the 1800RESPECT / 000 line that sits on every screen.',
    23: 'Recap the four key messages and note the support line is always present. Then open it up for discussion.',
}

GOLD = RGBColor(0xF0, 0xC6, 0x6A); INK = RGBColor(0x2A, 0x1F, 0x10)
DARK = RGBColor(0x20, 0x18, 0x0F); PANEL = RGBColor(0x17, 0x11, 0x0A)
CREAM = RGBColor(0xF5, 0xEF, 0xE2); MUTE = RGBColor(0xCF, 0xC7, 0xB8)
HARMLINE = RGBColor(0x8A, 0x94, 0x9A); SUBLINE = RGBColor(0x6B, 0x60, 0x50)
SERIF, SANS = 'Georgia', 'Helvetica'

KEY_MESSAGES = [
    'Violence is always a choice and caused by gender inequality',
    'Gender inequality, pressure and systems all matter',
    'Understanding causes helps prevent harm — it never excuses it',
    'Men can lead positive change',
]
SUPPORT = 'If this raised something for you: 1800RESPECT (1800 737 732) · in an emergency call 000'

prs = Presentation()
prs.slide_width = Emu(12192000)   # 13.333 in
prs.slide_height = Emu(6858000)   # 7.5 in (16:9)
blank = prs.slide_layouts[6]


def add_para(tf, text, size, font, color, *, italic=False, bold=False, align=PP_ALIGN.CENTER, before=0, after=6, spacing=None):
    p = tf.paragraphs[0] if (len(tf.paragraphs) == 1 and not tf.paragraphs[0].runs) else tf.add_paragraph()
    p.alignment = align
    if before: p.space_before = Pt(before)
    p.space_after = Pt(after)
    if spacing: p.line_spacing = spacing
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.name = font; r.font.italic = italic; r.font.bold = bold
    r.font.color.rgb = color
    return p


def panel(slide, x, y, w, h):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = PANEL
    shp.line.color.rgb = GOLD; shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp


def no_click_advance(slide):
    sld = slide._element
    tr = sld.makeelement(qn('p:transition'), {'advClick': '0'})
    cmo = sld.find(qn('p:clrMapOvr'))
    (cmo if cmo is not None else sld.find(qn('p:cSld'))).addnext(tr)


# ---- slide 0: title ----
t = prs.slides.add_slide(blank)
t.shapes.add_picture(HERO, 0, 0, width=prs.slide_width, height=prs.slide_height)
panel(t, Emu(2396000), Emu(1150000), Emu(7400000), Emu(4150000))
tb = t.shapes.add_textbox(Emu(2700000), Emu(1430000), Emu(6800000), Emu(3100000)).text_frame
tb.word_wrap = True
add_para(tb, 'INTERACTIVE  WALKTHROUGH', 13, SANS, GOLD, bold=True, after=14)
add_para(tb, 'Shaping Change', 40, SERIF, CREAM, italic=True, bold=True, after=6)
add_para(tb, 'Same roots, different trees.', 22, SERIF, GOLD, italic=True, after=14)
add_para(tb, 'Follow Orion: first watch the pressures grow a strained tree, then help '
             'him rebuild. The same pressures can grow very different trees — what changes '
             'the outcome is the beliefs he holds, how he responds, and whether he leads.',
         15, SANS, MUTE, after=0, spacing=1.2)
t.notes_slide.notes_text_frame.text = NOTES[0]

# ---- slides 1..22: screens ----
screen_slides = []
for k, name in enumerate(SCREENS):
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(os.path.join(IMG, name), 0, 0, width=prs.slide_width, height=prs.slide_height)
    s.notes_slide.notes_text_frame.text = NOTES[k + 1]
    screen_slides.append(s)

# ---- slide 23: closing / key messages ----
c = prs.slides.add_slide(blank)
c.shapes.add_picture(HERO, 0, 0, width=prs.slide_width, height=prs.slide_height)
panel(c, Emu(1500000), Emu(760000), Emu(9192000), Emu(5300000))
cb = c.shapes.add_textbox(Emu(1900000), Emu(1040000), Emu(8400000), Emu(3700000)).text_frame
cb.word_wrap = True
add_para(cb, 'KEY  MESSAGES', 13, SANS, GOLD, bold=True, align=PP_ALIGN.LEFT, after=12)
for m in KEY_MESSAGES:
    add_para(cb, '•   ' + m, 18, SANS, CREAM, align=PP_ALIGN.LEFT, after=10)
add_para(cb, SUPPORT, 13, SANS, MUTE, italic=True, align=PP_ALIGN.LEFT, before=8, after=0)
c.notes_slide.notes_text_frame.text = NOTES[CLOSING]

slides = [t] + screen_slides + [c]
for s in slides:
    no_click_advance(s)


# ---- buttons + click actions ----
def style(shape, kind):
    shape.fill.solid(); shape.line.width = Pt(1.25); shape.shadow.inherit = False
    if kind == 'go':
        shape.fill.fore_color.rgb = GOLD; shape.line.color.rgb = GOLD; return INK
    if kind == 'harm':
        shape.fill.fore_color.rgb = DARK; shape.line.color.rgb = HARMLINE; return CREAM
    shape.fill.fore_color.rgb = DARK; shape.line.color.rgb = SUBLINE; return MUTE


def button_geoms(idx, n):
    """Positions for n buttons on a slide. Screen slides keep buttons in the
    bottom-LEFT band so they never cover the game panel on the right."""
    if idx == TITLE:
        return [(Emu(4946000), Emu(4520000), Emu(2300000), Emu(620000))]
    if idx == CLOSING:
        xs = [Emu(2100000), Emu(6300000)]
        return [(xs[j], Emu(5200000), Emu(3500000), Emu(560000)) for j in range(n)]
    # screen slides: left band x in [300000, 8600000], clear of the right-hand panel
    x0, x_end, gap = 300000, 8600000, 240000
    h = 720000 if n >= 3 else 560000
    y = Emu(5830000) if n >= 3 else Emu(5990000)
    bw = min(4030000, int((x_end - x0 - gap * (n - 1)) / n))
    return [(Emu(x0 + j * (bw + gap)), y, Emu(bw), h) for j in range(n)]


def place_button(slide, geom, label, kind, target, small):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, *geom)
    txtcol = style(shp, kind)
    tf = shp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    r.font.size = Pt(13 if small else 16); r.font.bold = True; r.font.name = SANS
    r.font.color.rgb = txtcol
    shp.click_action.target_slide = slides[target]


for idx, btns in NAV.items():
    geoms = button_geoms(idx, len(btns))
    small = len(btns) >= 3
    for geom, (label, target, kind) in zip(geoms, btns):
        place_button(slides[idx], geom, label, kind, target, small)

prs.core_properties.title = 'Shaping Change — Clickable Walkthrough'
prs.core_properties.author = 'Shaping Change'

out = os.path.join(ROOT, 'docs', 'Power-Pressure-Choice-Deck.pptx')
prs.save(out)
print('saved', os.path.relpath(out, ROOT), '(%d slides)' % len(slides))
