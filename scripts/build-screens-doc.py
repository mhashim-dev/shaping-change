#!/usr/bin/env python3
"""Build a Word 'Screen Designs & Descriptions' reference — one screen per page,
each with its design mockup image, what's on screen, and its purpose.

    pip install python-docx
    python3 scripts/build-screens-doc.py   # -> docs/Shaping-Change-Screen-Designs.docx
"""
import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from screens_data import SCREENS, ENHANCEMENTS   # noqa: E402  (canonical per-screen + enhancements)

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMG = os.path.join(ROOT, 'docs', 'figma-screens')

BROWN = RGBColor(0x5A, 0x3E, 0x1B)
GOLDINK = RGBColor(0x8A, 0x6A, 0x1E)
GREY = RGBColor(0x70, 0x6A, 0x60)
INK = RGBColor(0x2A, 0x24, 0x1C)


doc = Document()
# margins
for s in doc.sections:
    s.left_margin = s.right_margin = Inches(1.0)
    s.top_margin = s.bottom_margin = Inches(0.9)

normal = doc.styles['Normal'].font
normal.name = 'Calibri'; normal.size = Pt(11)


def run(p, text, *, size=11, bold=False, italic=False, color=None, font='Calibri'):
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic; r.font.name = font
    if color is not None: r.font.color.rgb = color
    return r


# ---- title block ----
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
run(p, 'Shaping Change', size=26, bold=True, color=BROWN, font='Georgia')
p2 = doc.add_paragraph(); run(p2, 'Screen Designs & Descriptions', size=16, italic=True, color=GOLDINK, font='Georgia')
p3 = doc.add_paragraph()
run(p3, 'Building strong roots for safety  ·  24 screens  ·  Edu & PVAW–approved content', size=10.5, color=GREY)
p3b = doc.add_paragraph()
run(p3b, 'Updated July 2026 — the approved content is unchanged; a final section covers new '
         'accessibility, community-language, take-home, facilitator and offline enhancements for your review.',
    size=10.5, italic=True, color=GOLDINK)
doc.add_paragraph()
intro = doc.add_paragraph()
run(intro, 'This is a screen-by-screen reference for the activity — one screen per page, with its design '
           'mockup, what appears on screen, and its purpose. The story runs in two phases: in ', size=11)
run(intro, 'Phase 1', bold=True)
run(intro, ' the player experiences the problem and the tree declines; in ', size=11)
run(intro, 'Phase 2', bold=True)
run(intro, ' the player helps Orion rebuild, and the tree heals (or stays strained). The tree carries '
           'the meaning throughout — roots are pressures, soil is beliefs, the trunk is behaviour, and the '
           'branches are the impact on family. The 1800RESPECT / 000 support line appears on every screen.',
    size=11)

# ---- one screen per page ----
for i, (f, title, phase, on_screen, purpose) in enumerate(SCREENS):
    doc.add_page_break()
    h = doc.add_paragraph()
    run(h, f'Screen {f[:2]}', size=12, bold=True, color=GOLDINK)
    run(h, f'   ·   {title}', size=16, bold=True, color=BROWN, font='Georgia')
    ph = doc.add_paragraph(); run(ph, phase, size=10, italic=True, color=GREY)
    img = os.path.join(IMG, f + '.png')
    doc.add_picture(img, width=Inches(6.5))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(cap, f + '.png', size=8.5, italic=True, color=GREY)
    a = doc.add_paragraph(); run(a, 'On screen   ', size=11, bold=True, color=INK); run(a, on_screen, size=11)
    b = doc.add_paragraph(); run(b, 'Purpose   ', size=11, bold=True, color=INK); run(b, purpose, size=11)

# ---- enhancements section (for team / Edu / PVAW review) ----
GREEN = RGBColor(0x2E, 0x6B, 0x3A)
BLUE = RGBColor(0x1F, 0x4E, 0x79)

ENH_IMG = os.path.join(ROOT, 'docs', 'images', 'enhancements')

doc.add_page_break()
sh = doc.add_paragraph()
run(sh, 'Enhancements since content sign-off', size=20, bold=True, color=BROWN, font='Georgia')
sp = doc.add_paragraph()
run(sp, 'For your review — July 2026', size=12, italic=True, color=GOLDINK, font='Georgia')
si = doc.add_paragraph()
run(si, 'The approved content and the two-phase story are unchanged. The additions below improve '
        'accessibility, reach and delivery. A few add small optional elements or presentation improvements; '
        'none change the approved wording of the scenario. Items marked ', size=11)
run(si, 'For your sign-off', size=11, bold=True, color=GREEN)
run(si, ' are where we would value Edu or PVAW’s confirmation.', size=11)

for title, audience, images, paras, flag in ENHANCEMENTS:
    doc.add_page_break()
    h = doc.add_paragraph(); run(h, title, size=15, bold=True, color=BROWN, font='Georgia')
    ta = doc.add_paragraph(); run(ta, audience.upper(), size=9.5, bold=True, color=BLUE)
    for para in paras:
        pp = doc.add_paragraph(); run(pp, para, size=11)
    if flag:
        fp = doc.add_paragraph(); run(fp, '✔ ' + flag, size=11, bold=True, color=GREEN)
    if images:
        for imgname in images:
            path = os.path.join(ENH_IMG, imgname)
            if os.path.exists(path):
                doc.add_picture(path, width=Inches(6.2))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

out = os.path.join(ROOT, 'docs', 'Shaping-Change-Screen-Designs.docx')
doc.core_properties.title = 'Shaping Change — Screen Designs & Descriptions'
doc.core_properties.author = 'Shaping Change'
doc.save(out)
print('saved', os.path.relpath(out, ROOT), f'({len(SCREENS)} screens)')
