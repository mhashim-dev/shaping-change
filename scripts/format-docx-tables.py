#!/usr/bin/env python3
"""Make the tables in a pandoc-generated .docx clean and readable, in place:
  - column widths fitted to each column's content (so text columns aren't cramped)
  - a shaded, bold header row
  - subtle alternating row banding
  - thin grey grid borders + comfortable cell padding + 10pt text

Usage:  PYTHONPATH=/tmp/pptxlib python3 scripts/format-docx-tables.py <file.docx> [more.docx ...]
"""
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HEADER_FILL = '34495E'   # dark slate
BAND_FILL = 'EEF2F5'     # very light blue-grey
BORDER = 'C3CCD3'
HEADER_TEXT = RGBColor(0xFF, 0xFF, 0xFF)
BODY_TEXT = RGBColor(0x22, 0x28, 0x2D)
FONT = 'Calibri'
HEADER_PT, BODY_PT = 10, 10
MIN_COL_IN = 0.75        # never let a column get narrower than this


def _set(el, tag, **attrs):
    sub = OxmlElement(tag)
    for k, v in attrs.items():
        sub.set(qn(k), v)
    el.append(sub)
    return sub


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('w:shd')):
        tcPr.remove(old)
    _set(tcPr, 'w:shd', **{'w:val': 'clear', 'w:color': 'auto', 'w:fill': fill})


def set_borders(table):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(old)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        _set(borders, 'w:' + edge, **{'w:val': 'single', 'w:sz': '4', 'w:space': '0', 'w:color': BORDER})
    tblPr.append(borders)


def cell_margins(table, top=40, bottom=40, left=90, right=90):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn('w:tblCellMar')):
        tblPr.remove(old)
    mar = OxmlElement('w:tblCellMar')
    for side, val in (('top', top), ('bottom', bottom), ('left', left), ('right', right)):
        _set(mar, 'w:' + side, **{'w:w': str(val), 'w:type': 'dxa'})
    tblPr.append(mar)


def fixed_layout(table):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn('w:tblLayout')):
        tblPr.remove(old)
    _set(tblPr, 'w:tblLayout', **{'w:type': 'fixed'})


def style_run_text(cell, color, bold, size_pt):
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        for r in p.runs:
            r.font.name = FONT
            r.font.size = Pt(size_pt)
            r.font.color.rgb = color
            if bold is not None:
                r.font.bold = bold


def fit_widths(table, usable_in):
    rows = table.rows
    ncol = len(table.columns)
    # average content length per column → proportional widths, with a floor
    totals = [0] * ncol
    for r in rows:
        for j, c in enumerate(r.cells[:ncol]):
            totals[j] += max(1, len(c.text))
    grand = sum(totals) or ncol
    raw = [usable_in * t / grand for t in totals]
    # apply floor, then renormalise the remainder
    widths = [max(MIN_COL_IN, w) for w in raw]
    over = sum(widths) - usable_in
    if over > 0:
        slack = [w - MIN_COL_IN for w in widths]
        s = sum(slack) or 1
        widths = [w - over * (sl / s) for w, sl in zip(widths, slack)]
    # write grid + every cell (fixed layout needs per-cell widths)
    grid = table._tbl.find(qn('w:tblGrid'))
    if grid is not None:
        for gc, w in zip(grid.findall(qn('w:gridCol')), widths):
            gc.set(qn('w:w'), str(int(w * 1440)))
    for r in rows:
        for j, c in enumerate(r.cells[:ncol]):
            c.width = Inches(widths[j])


def format_table(table, usable_in):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.allow_autofit = False
    fixed_layout(table)
    set_borders(table)
    cell_margins(table)
    fit_widths(table, usable_in)
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            if i == 0:
                shade(cell, HEADER_FILL)
                style_run_text(cell, HEADER_TEXT, True, HEADER_PT)
            else:
                if i % 2 == 0:
                    shade(cell, BAND_FILL)
                style_run_text(cell, BODY_TEXT, None, BODY_PT)


def process(path):
    doc = Document(path)
    sec = doc.sections[0]
    pw, lm, rm = sec.page_width, sec.left_margin, sec.right_margin
    if pw is None or lm is None or rm is None:
        usable = 6.5  # Letter (8.5in) with 1in margins — pandoc default
    else:
        usable = (pw - lm - rm) / 914400.0  # EMU -> in
    for t in doc.tables:
        format_table(t, usable)
    doc.save(path)
    print('formatted %d table(s) in %s (usable width %.2f in)' % (len(doc.tables), path, usable))


if __name__ == '__main__':
    for p in sys.argv[1:]:
        process(p)
