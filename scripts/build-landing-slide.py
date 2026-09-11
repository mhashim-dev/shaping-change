#!/usr/bin/env python3
"""Build a single review slide for the proposed LANDING-PAGE intro (before the game).
Styled to match the Screen Designs deck so it can be dropped straight in.

    pip install python-pptx
    python3 scripts/build-landing-slide.py   # -> docs/Shaping-Change-Landing-Page-Slide.pptx
"""
import os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GOLD = RGBColor(0xF0, 0xC6, 0x6A); INK = RGBColor(0x2A, 0x1F, 0x10)
DARK = RGBColor(0x1B, 0x14, 0x0C); PANEL = RGBColor(0x17, 0x11, 0x0A)
CREAM = RGBColor(0xF5, 0xEF, 0xE2); MUTE = RGBColor(0xCF, 0xC7, 0xB8)
LINEC = RGBColor(0x4A, 0x3E, 0x2A)
SERIF, SANS = 'Georgia', 'Helvetica'
W, H = 12192000, 6858000

prs = Presentation(); prs.slide_width = Emu(W); prs.slide_height = Emu(H)
s = prs.slides.add_slide(prs.slide_layouts[6])


def rect(shape_type, x, y, w, h, fill, line=None, lw=1.0):
    sp = s.shapes.add_shape(shape_type, Emu(x), Emu(y), Emu(w), Emu(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb = line; sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    return sp


def tb(x, y, w, h):
    t = s.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h)).text_frame
    t.word_wrap = True
    return t


def para(tf, runs, *, align=PP_ALIGN.LEFT, before=0, after=6, spacing=1.1, reuse=False):
    p = tf.paragraphs[0] if (reuse and len(tf.paragraphs) == 1 and not tf.paragraphs[0].runs) else tf.add_paragraph()
    p.alignment = align
    if before: p.space_before = Pt(before)
    p.space_after = Pt(after); p.line_spacing = spacing
    for (text, size, font, color, bold, italic) in runs:
        r = p.add_run(); r.text = text
        r.font.size = Pt(size); r.font.name = font; r.font.bold = bold; r.font.italic = italic
        r.font.color.rgb = color
    return p


# background + header
rect(MSO_SHAPE.RECTANGLE, 0, 0, W, H, DARK)
eb = tb(520000, 300000, 11000000, 460000)
para(eb, [('SHAPING CHANGE    ·    LANDING PAGE  (BEFORE THE GAME)    ·    FOR REVIEW', 12, SANS, GOLD, True, False)], after=0, reuse=True)
rect(MSO_SHAPE.RECTANGLE, 520000, 740000, W - 1040000, 9000, LINEC)
ti = tb(520000, 850000, 11150000, 560000)
para(ti, [('A short intro on the landing page — the game itself still opens with the character intro',
           15, SERIF, CREAM, False, True)], after=0, reuse=True)

# ---- the landing-page "card" (dark panel, gold border) ----
CX, CY, CW, CH = 760000, 1560000, 10672000, 5060000
card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(CX), Emu(CY), Emu(CW), Emu(CH))
card.fill.solid(); card.fill.fore_color.rgb = PANEL; card.line.color.rgb = GOLD; card.line.width = Pt(1.25)
card.shadow.inherit = False

body = tb(CX + 420000, CY + 300000, CW - 840000, 3450000)
para(body, [('Shaping Change', 25, SERIF, CREAM, True, True)], after=2, reuse=True)
para(body, [('Building strong roots for safety', 14, SERIF, GOLD, False, True)], after=12)
para(body, [('Welcome. Shaping Change is a short, reflective activity about building safe and respectful homes.',
             13.5, SANS, CREAM, False, False)], after=8, spacing=1.14)
para(body, [('Starting again in a new country changes many things — work, language, culture, and roles at home. '
             'It can bring hope, and also real stress. This activity looks at how those pressures shape the way we '
             'feel and act — and how, as leaders in our families and communities, we can respond in healthier ways.',
             13.5, SANS, MUTE, False, False)], after=8, spacing=1.14)
para(body, [('You’ll follow one person’s story and make a few choices. A tree grows alongside the story to show how '
             'pressures, beliefs and choices shape family life. There are no right or wrong answers, and no blame — '
             'it’s about understanding patterns, and the choices we can make.',
             13.5, SANS, MUTE, False, False)], after=0, spacing=1.14)

expect = tb(CX + 420000, CY + 3520000, CW - 840000, 900000)
para(expect, [('WHAT TO EXPECT', 11, SANS, GOLD, True, False)], after=4, reuse=True)
para(expect, [('About 10–15 minutes, at your own pace     ·     Simple English     ·     '
               'Works on phone, tablet or computer     ·     Private — your answers stay on your device',
               12.5, SANS, CREAM, False, False)], after=0, spacing=1.12)

# CTA button + support footer
btn = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(CX + 420000), Emu(CY + 4180000), Emu(2650000), Emu(520000))
btn.fill.solid(); btn.fill.fore_color.rgb = GOLD; btn.line.fill.background(); btn.shadow.inherit = False
bt = btn.text_frame; bt.word_wrap = True
para(bt, [('Start the activity  →', 14, SANS, INK, True, False)], align=PP_ALIGN.CENTER, after=0, reuse=True)

sup = tb(CX + 3250000, CY + 4230000, CW - 3700000, 460000)
para(sup, [('Support is always available — 1800RESPECT (1800 737 732), or 000 in an emergency.',
            11.5, SANS, MUTE, False, True)], after=0, reuse=True)

s.notes_slide.notes_text_frame.text = (
    'Proposed intro for the landing page — for your review.\n\n'
    'Where it sits: on the page that links to the game. Agreed with Ali & Anu (28 Jul) that the short '
    'intro sits here, and the game itself still opens with the character intro (Meet Orion).\n\n'
    'Design choices to note: no statistics here (figures stay at the end, per PVAW guidance); warm, '
    'non-judgemental, strengths-based; simple English for the EAL audience; and the 1800RESPECT / 000 '
    'support line is shown on the landing page too.\n\n'
    'For your review: does the tone and wording feel right? Full copy (plus a shorter blurb and an even '
    'shorter version) is in docs/LANDING_PAGE_INTRO.md.')

out = os.path.join(ROOT, 'docs', 'Shaping-Change-Landing-Page-Slide.pptx')
prs.core_properties.title = 'Shaping Change — Landing Page Intro (for review)'
prs.core_properties.author = 'Shaping Change'
prs.save(out)
print('saved', os.path.relpath(out, ROOT), '(%d slide)' % len(prs.slides._sldIdLst))
