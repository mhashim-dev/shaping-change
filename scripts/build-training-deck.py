#!/usr/bin/env python3
"""Build the team TRAINING deck (to present from during the handover session).
Concise slides + speaker notes for a non-technical presenter.
    python3 scripts/build-training-deck.py  -> docs/Shaping-Change-Team-Training.pptx
"""
import os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GOLD = RGBColor(0xF0, 0xC6, 0x6A); DARK = RGBColor(0x1B, 0x14, 0x0C); PANEL = RGBColor(0x17, 0x11, 0x0A)
CREAM = RGBColor(0xF5, 0xEF, 0xE2); MUTE = RGBColor(0xCF, 0xC7, 0xB8); LINEC = RGBColor(0x4A, 0x3E, 0x2A)
GREEN = RGBColor(0x8F, 0xD1, 0x9A); BLUE = RGBColor(0x8F, 0xC0, 0xFF); CODEC = RGBColor(0xE9, 0xE2, 0xD2)
SERIF, SANS, MONO = 'Georgia', 'Helvetica', 'Consolas'

W, H = 12192000, 6858000
prs = Presentation(); prs.slide_width = Emu(W); prs.slide_height = Emu(H)
blank = prs.slide_layouts[6]


def rect(slide, x, y, w, h, fill, rounded=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                 Emu(x), Emu(y), Emu(w), Emu(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill; shp.line.fill.background(); shp.shadow.inherit = False
    return shp


def tb(slide, x, y, w, h):
    t = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h)).text_frame; t.word_wrap = True
    return t


def para(tf, runs, *, align=PP_ALIGN.LEFT, before=0, after=6, spacing=1.08, reuse=False):
    p = tf.paragraphs[0] if (reuse and len(tf.paragraphs) == 1 and not tf.paragraphs[0].runs) else tf.add_paragraph()
    p.alignment = align
    if before: p.space_before = Pt(before)
    p.space_after = Pt(after); p.line_spacing = spacing
    if isinstance(runs, str): runs = [(runs, 18, SANS, CREAM, False, False)]
    for (text, size, font, color, bold, italic) in runs:
        r = p.add_run(); r.text = text
        r.font.size = Pt(size); r.font.name = font; r.font.bold = bold; r.font.italic = italic; r.font.color.rgb = color
    return p


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def base(kicker='SHAPING CHANGE   ·   TEAM TRAINING'):
    s = prs.slides.add_slide(blank); rect(s, 0, 0, W, H, DARK)
    eb = tb(s, 560000, 330000, 10000000, 460000)
    para(eb, [(kicker, 12, SANS, GOLD, True, False)], after=0, reuse=True)
    rect(s, 560000, 780000, W - 1120000, 9000, LINEC)
    return s


def title_body(s, title, sub=None):
    t = tb(s, 560000, 1080000, W - 1120000, 1300000)
    para(t, [(title, 30, SERIF, CREAM, True, True)], after=(4 if sub else 10), reuse=True)
    if sub: para(t, [(sub, 15, SANS, GOLD, False, True)], after=0)
    return t


def bullets(s, items, y=2350000, size=17, gap=11):
    t = tb(s, 560000, y, W - 1120000, H - y - 500000)
    first = True
    for it in items:
        if isinstance(it, tuple):
            head, body = it
            para(t, [('•  ', size, SANS, GOLD, True, False), (head + '  ', size, SANS, CREAM, True, False),
                     (body, size, SANS, MUTE, False, False)], after=gap, before=(0 if first else 0), spacing=1.1)
        else:
            para(t, [('•  ', size, SANS, GOLD, True, False), (it, size, SANS, CREAM, False, False)], after=gap, spacing=1.1)
        first = False
    return t


def codebox(s, x, y, w, h, lines):
    p = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    p.fill.solid(); p.fill.fore_color.rgb = PANEL; p.line.color.rgb = LINEC; p.line.width = Pt(1); p.shadow.inherit = False
    t = tb(s, x + 220000, y + 160000, w - 440000, h - 320000)
    for i, ln in enumerate(lines):
        para(t, [(ln if ln else ' ', 13, MONO, CODEC, False, False)], after=2, spacing=1.05, reuse=(i == 0))
    return p


# ---------------- 1 COVER ----------------
c = prs.slides.add_slide(blank); rect(c, 0, 0, W, H, PANEL)
rect(c, 0, 0, 150000, H, GOLD)
ct = tb(c, 900000, 2000000, 10000000, 3200000)
para(ct, [('TEAM TRAINING  &  HANDOVER', 15, SANS, GOLD, True, False)], after=16, reuse=True)
para(ct, [('Shaping Change', 46, SERIF, CREAM, True, True)], after=6)
para(ct, [('Maintaining and updating the game — a hands-on session', 20, SERIF, GOLD, False, True)], after=20)
para(ct, [('No coding background needed  ·  [Presenter]  ·  [Date]', 15, SANS, MUTE, False, False)], after=0)
notes(c, 'Welcome. The goal of today is simple: by the end, you will be able to make changes to the game '
         'yourselves. You do NOT need a coding background — most of the work is editing a list of words and '
         'choices. We will do a real change together, live. Keep it relaxed and hands-on.')

# ---------------- 2 AGENDA ----------------
s = base(); title_body(s, 'What we’ll cover')
bullets(s, ['How the game is built — the simple version',
            'Who does what on the team',
            'The one file you’ll edit most',
            'Making a real change together, start to finish',
            'The safety check, and how to share the game'])
notes(s, 'Set expectations: this is practical, not theory. The single most useful moment will be when we make a '
         'change together and watch it appear. Everything else supports that.')

# ---------------- 3 WHO DOES WHAT ----------------
s = base(); title_body(s, 'Who does what on the team', 'Most updates need no coding at all')
pw = 5300000; px1 = 560000; px2 = 560000 + pw + 520000; py = 2450000; ph = 3250000
for (x, tag, tagc, ttl, body) in [
    (px1, 'ANYONE ON THE TEAM', GREEN, 'The content',
     'Edit the words on any screen · change or add an answer option · change the support links · add new '
     'screens. This is careful list-editing in one file — not programming.'),
    (px2, 'YOUR HTML / CSS PERSON', BLUE, 'The technical lead',
     'Install the tools · run the short commands · rebuild the shareable file · handle the “look” (colours, '
     'fonts, spacing — that’s CSS). First person to call if something breaks.')]:
    pnl = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(py), Emu(pw), Emu(ph))
    pnl.fill.solid(); pnl.fill.fore_color.rgb = PANEL; pnl.line.color.rgb = LINEC; pnl.line.width = Pt(1); pnl.shadow.inherit = False
    pt = tb(s, x + 320000, py + 280000, pw - 640000, ph - 560000)
    para(pt, [(tag, 13, SANS, tagc, True, False)], after=6, reuse=True)
    para(pt, [(ttl, 23, SERIF, CREAM, True, True)], after=12)
    para(pt, [(body, 15.5, SANS, MUTE, False, False)], after=0, spacing=1.16)
notes(s, 'This is the reassuring slide. Point at the left panel: this is where 90% of the work is, and anyone can '
         'do it. The right panel is for our one HTML/CSS person, who becomes the technical lead — they run the '
         'commands and own the look.')

# ---------------- 4 BIG PICTURE ----------------
s = base(); title_body(s, 'The big picture', 'There are two copies of the same game')
bullets(s, [('The dev app —', 'what you edit and preview live while working (runs on your computer only).'),
            ('The offline game —', 'the single file you SHARE: “Power-Pressure-Choice-POC.html”. Runs on any device, no internet.'),
            ('Both read the same words & choices —', 'so you write the content once. The dev app updates instantly; the offline game is “rebuilt” with one command when you’re ready to share.')])
notes(s, 'Make the two-copies distinction clear, because it explains why we sometimes “rebuild”. The offline file '
         'is the thing we hand to people — it needs no internet and keeps everything on the player’s device.')

# ---------------- 5 KEY IDEA ----------------
s = base(); title_body(s, 'The one idea to remember')
t = tb(s, 560000, 2300000, W - 1120000, 3200000)
para(t, [('The words & choices are kept separate from the drawing & mechanics.', 22, SERIF, CREAM, True, True)], after=18, reuse=True)
para(t, [('So about ', 18, SANS, CREAM, False, False), ('90% of changes live in ONE plain file:', 18, SANS, CREAM, True, False)], after=8)
codebox(s, 560000, 3450000, 7200000, 620000, ['lib/game-content.js'])
para(tb(s, 560000, 4300000, W - 1120000, 900000),
     [('If you can edit a list, you can edit the game.', 19, SANS, GOLD, False, True)], after=0, reuse=True)
notes(s, 'Hammer this point. The scary-sounding technology (Next.js) draws the tree and runs the page — but the '
         'part you touch is a simple list of screens. Repeat: if you can edit a list, you can edit the game.')

# ---------------- 6 WORDS YOU'LL HEAR ----------------
s = base(); title_body(s, 'A few words you’ll hear')
bullets(s, [('Next.js / React —', 'the engine that turns our files into a web page. You rarely touch it.'),
            ('Node.js / npm —', 'a free program that runs the short commands. Installed once.'),
            ('Terminal —', 'the text window where you type commands (in the editor: View → Terminal).'),
            ('localhost:3000 —', 'your own private preview of the game while you work.'),
            ('VS Code —', 'the free editor we use to open the project and edit files.')], size=16, gap=10)
notes(s, 'Don’t dwell — just demystify the words so nobody freezes when they hear them. You’ll see each one in '
         'action in a minute.')

# ---------------- 7 SETUP ----------------
s = base(); title_body(s, 'One-time setup', 'WHO: your HTML/CSS person, once per computer')
bullets(s, ['Install Node.js — nodejs.org, the “LTS” version, click through the defaults',
            'Install VS Code — code.visualstudio.com',
            'In VS Code: File → Open Folder → the “ames-tree-main” folder',
            'Open the Terminal (View → Terminal) and run:'], y=2350000, gap=9)
codebox(s, 560000, 4650000, 5200000, 620000, ['npm install'])
notes(s, 'This is a one-time thing per computer. It downloads the project’s parts. After this, the everyday '
         'commands just work.')

# ---------------- 8 SEE CHANGES ----------------
s = base(); title_body(s, 'See the game and your changes', 'WHO: anyone')
bullets(s, ['In the Terminal, run this and leave it running:'], y=2300000, gap=6)
codebox(s, 560000, 2850000, 5200000, 560000, ['npm run dev'])
bullets(s, ['Open a browser to  http://localhost:3000',
            'Edit a file and Save → the page updates by itself',
            'The offline game: just double-click the HTML file'], y=3700000, gap=9)
notes(s, 'Show this live if you can: run npm run dev, open localhost:3000, and leave it up for the rest of the '
         'session so every change we make shows instantly.')

# ---------------- 9 FILE MAP ----------------
s = base(); title_body(s, 'Which file does what')
bullets(s, [('lib/game-content.js —', 'every screen: the words, the choices, the order. YOUR MAIN FILE.'),
            ('app/globals.css —', 'colours, fonts, spacing (the “look”) — for the CSS person.'),
            ('validate-content.mjs —', 'the safety check. You run it; you don’t edit it.'),
            ('build-poc.cjs —', 'builds the offline shareable game. You run it; you don’t edit it.'),
            ('lib/tree-engine.js —', 'draws the tree and Orion. Advanced — leave it alone.')], size=16, gap=10)
notes(s, 'The takeaway: one file to know (game-content.js). Two scripts you RUN but never edit. And the drawing '
         'engine, which we simply don’t touch.')

# ---------------- 10 ONE SCREEN ----------------
s = base(); title_body(s, 'How one screen works')
para(tb(s, 560000, 1780000, W - 1120000, 500000),
     [('Each screen is one block in the list. Here is a real one:', 15, SANS, MUTE, False, True)], after=0, reuse=True)
codebox(s, 560000, 2300000, W - 1120000, 3550000, [
    "{",
    "  id: 'pressure_money',        // the screen's internal name",
    "  type: 'story',               // story = narration + a Continue button",
    "  prompt: 'Money worries',     // the big heading",
    "  hint: 'Back home, Orion ...',// the body text the player reads",
    "  btn: 'Continue',             // the button label",
    "  fx: { health: -7 },          // effect on the hidden 'health' score",
    "  next: 'pressure_role'        // which screen comes next (by its id)",
    "}"])
notes(s, 'Walk through the fields slowly. The ones people will actually edit: prompt (heading), hint (body text), '
         'btn (button). “next” is how screens are chained together. Choice screens have an “options” list instead '
         'of a single “next” — same idea, one per button.')

# ---------------- 11 HOW IT CONNECTS ----------------
s = base(); title_body(s, 'How it all connects', 'and how the ending is chosen')
bullets(s, [('“next” —', 'links each screen to the one after it, by its id. Follow the “next” values to trace the whole story.'),
            ('the hidden “health” score —', 'starts at 68; each choice nudges it up or down — that’s what makes the tree greener or barer.'),
            ('the ending —', 'the game reads the score: 70+ = healthy tree · 40–69 = healing · below 40 = under strain.')])
notes(s, 'Keep this light. People don’t need the maths — just the idea that healthy choices lift the score and a '
         'higher score = a healthier tree and ending.')

# ---------------- 12 RECIPES ----------------
s = base(); title_body(s, 'The common changes (recipes)', 'each is a step-by-step in the guide')
bullets(s, ['Change the words on a screen',
            'Change or add an answer option',
            'Change the support links (last screen)',
            'Add a whole new screen',
            'Change the picture (thought bubble) on a screen'])
notes(s, 'These are the everyday jobs, and the written guide has numbered steps for each. Don’t read them out — '
         'just show that the recipes exist, then we’ll actually do the first one together.')

# ---------------- 13 LIVE DEMO ----------------
s = prs.slides.add_slide(blank); rect(s, 0, 0, W, H, PANEL); rect(s, 0, 0, W, 150000, GOLD)
dt = tb(s, 700000, 900000, W - 1400000, 1100000)
para(dt, [('LET’S DO IT LIVE', 15, SANS, GOLD, True, False)], after=8, reuse=True)
para(dt, [('Change a screen together', 34, SERIF, CREAM, True, True)], after=0)
bullets(s, ['1.  Find the screen in  lib/game-content.js',
            '2.  Edit the words, and Save',
            '3.  Run:  npm run test:content   (the safety check)',
            '4.  Watch it appear at  localhost:3000',
            '5.  Run:  npm run build:poc   (updates the file we share)'], y=2350000, size=18, gap=13)
notes(s, 'THIS IS THE HEART OF THE SESSION. Pick something small and visible — e.g. change the “Money worries” '
         'heading or a sentence. Go slowly, narrate each step, and let people follow along on their own machines '
         'if they can. When the change appears in the browser, that’s the “aha” moment — pause on it.')

# ---------------- 14 SAFETY CHECK ----------------
s = base(); title_body(s, 'The safety check', 'run it after every change')
codebox(s, 560000, 1900000, 5600000, 560000, ['npm run test:content'])
bullets(s, ['Because this is a violence-PREVENTION tool, some rules must never break by accident. It confirms:',
            ('', 'a harmful choice can never heal the tree · help-seeking always heals'),
            ('', 'the 1800RESPECT / 000 support line is on every screen'),
            ('', 'all three outcomes stay reachable · every link is real · statistics keep their sources')], y=2750000, size=15, gap=9)
para(tb(s, 560000, 5650000, W - 1120000, 700000),
     [('“✓ All 10 invariant groups hold” = you’re good. An error names exactly what to fix.', 15, SANS, GOLD, False, True)], after=0, reuse=True)
notes(s, 'Stress this: always run it before sharing. It’s our safety net that protects the Edu/PVAW sign-offs. If '
         'it flags something, it tells you the exact rule and screen — it’s helpful, not scary.')

# ---------------- 15 WORKFLOW ----------------
s = base(); title_body(s, 'The everyday workflow')
codebox(s, 560000, 2100000, W - 1120000, 2400000, [
    "1.  Edit  lib/game-content.js   and Save",
    "2.  npm run test:content         (safety check — fix anything it flags)",
    "3.  Look at  http://localhost:3000",
    "4.  npm run build:poc            (updates the offline shareable game)",
    "5.  Share  Power-Pressure-Choice-POC.html"])
notes(s, 'This is the loop to memorise. Steps 1–3 are for anyone; steps 4–5 (the rebuild + share) are a natural '
         'fit for the technical lead, though the commands are one line each.')

# ---------------- 16 CAREFUL ----------------
s = base(); title_body(s, 'A few things to be careful with')
bullets(s, ['Don’t rename a screen’s id unless you also fix every “next” that points to it',
            'Mind your commas and quotes — a blank page usually means a typo; just undo (Ctrl/Cmd + Z)',
            'Never remove the support line; never let a harmful choice heal the tree',
            'Any change to safety wording or statistics → back to Edu / PVAW before publishing'])
notes(s, 'Short and firm. The first two are the common beginner slips; the safety check catches most of them. The '
         'last two are non-negotiable pedagogy rules.')

# ---------------- 17 CHEAT SHEET / CLOSE ----------------
s = base(); title_body(s, 'Cheat-sheet & where to start')
codebox(s, 560000, 1950000, W - 1120000, 1850000, [
    "npm install            one-time setup on a new computer",
    "npm run dev            start the live preview (localhost:3000)",
    "npm run test:content   the pedagogy safety check (after every edit)",
    "npm run build:poc      rebuild the offline shareable game"])
bullets(s, [('Two files to remember:', 'lib/game-content.js (all the words & choices) and validate-content.mjs (the rules).'),
            ('You also have:', 'a full written guide (Word) and this deck to refer back to.')], y=4150000, size=15, gap=10)
para(tb(s, 560000, 5750000, W - 1120000, 600000), [('Thank you — questions & a go at it yourselves.', 18, SERIF, GOLD, False, True)], after=0, reuse=True)
notes(s, 'Close by handing it over: point them to the written guide for the step-by-step recipes, and encourage '
         'everyone to try one change themselves now while you’re in the room to help.')

out = os.path.join(ROOT, 'docs', 'Shaping-Change-Team-Training.pptx')
prs.core_properties.title = 'Shaping Change — Team Training & Handover'
prs.core_properties.author = 'Shaping Change'
prs.save(out)
print('saved', os.path.relpath(out, ROOT), '(%d slides)' % len(prs.slides._sldIdLst))
