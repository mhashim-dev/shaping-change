#!/usr/bin/env python3
"""Build a plain-English change log (Word) of the PAG-feedback-round updates, for sharing back
to the advisory group. Pulls the change list from screens_data.CHANGES so it can't drift from
the deck's "What changed" slides.
    python3 scripts/build-change-summary.py  -> docs/Shaping-Change-Changes-PAG-Round.docx
"""
import os, sys
from docx import Document
from docx.shared import Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from screens_data import CHANGES

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GOLD = RGBColor(0xB0, 0x86, 0x1a); DARK = RGBColor(0x2A, 0x1F, 0x10)
GREEN = RGBColor(0x1e, 0x7a, 0x3c); GREY = RGBColor(0x55, 0x4d, 0x40)

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10.5)


def heading(text, size, color=DARK, bold=True, before=12, after=4, italic=False):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size); r.font.color.rgb = color
    return p


def rule():
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(6)
    p.add_run('_' * 60).font.color.rgb = RGBColor(0xCC, 0xC4, 0xB6)


def bullet(text, source):
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(4)
    p.add_run(text + '  ')
    r = p.add_run('(' + source + ')'); r.italic = True; r.font.size = Pt(9.5); r.font.color.rgb = GREY


# header
heading('Shaping Change — Changes from the PAG Feedback', 17, DARK, before=0, after=2)
sub = doc.add_paragraph(); sub.paragraph_format.space_after = Pt(6)
rs = sub.add_run('Actioned August 2026, in response to the PAG feedback document (the “Note to designer” '
                 'section and Ali’s comments) and the PowerPoint review. Each change notes its source.')
rs.italic = True; rs.font.size = Pt(10); rs.font.color.rgb = GREY
rule()

for hi, (section, items) in enumerate(CHANGES):
    heading(section, 13, GOLD, before=(2 if hi == 0 else 12))
    for text, source in items:
        bullet(text, source)

# verification note
rule()
heading('Verified', 13, GREEN, before=2)
for line in [
    'The pedagogy safeguard check passes — all invariants hold across the 23 screens '
    '(all three outcomes stay reachable; a harmful choice can never heal the tree; the support line is on every screen).',
    'A full run-through of the game was tested end to end — no errors, and the story flows correctly.',
    'All 23 screen mockups were re-rendered, and the Word doc and the screen-designs deck were rebuilt to match.',
]:
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(4); p.add_run(line)

foot = doc.add_paragraph(); foot.paragraph_format.space_before = Pt(10)
rf = foot.add_run('The items under “Still to come” need illustration work or are planned for a future version, '
                  'and are flagged rather than built in this round.')
rf.italic = True; rf.font.size = Pt(9.5); rf.font.color.rgb = GREY

out = os.path.join(ROOT, 'docs', 'Shaping-Change-Changes-PAG-Round.docx')
doc.save(out)
print('saved', os.path.relpath(out, ROOT), '(%d change items)' % sum(len(i) for _, i in CHANGES))
