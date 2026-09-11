#!/usr/bin/env python3
"""Build the PAG feedback tracker (Word + Markdown) from the consolidated advisory-group
feedback: the Word doc body, its 6 margin comments (Ali/Anu), the 3 PowerPoint comments,
and Anu's covering email.  Single source of truth -> docs/*.

    python3 scripts/build-pag-feedback-tracker.py
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# Status vocabulary + colour
DONE='Done'; CODE='To do — code'; DESIGN='To do — design'; CONTENT='To do — content (PVAW)'
DECIDE='Needs decision'; FUTURE='Future version'; PROCESS='Process'
COLOR={DONE:RGBColor(0x1e,0x7a,0x3c), CODE:RGBColor(0xB0,0x86,0x1a), DESIGN:RGBColor(0x8a,0x5a,0x00),
       CONTENT:RGBColor(0x9c,0x27,0x10), DECIDE:RGBColor(0xb0,0x1a,0x1a), FUTURE:RGBColor(0x55,0x55,0x55),
       PROCESS:RGBColor(0x40,0x40,0x40)}

# theme -> list of (id, feedback, source, action/decision, status, owner)
THEMES = [
 ('A. Cross-cutting / framing', [
  ('A1','Keep the focus on Orion; remove references to "external action / the player as a community leader."',
        'Word comment 1 & 4; PPT (Anu)','Reframe to Orion-only for this version — drop the player-as-leader angle in copy and structure.',CODE,'Dev + Content'),
  ('A2','Remove the optional readiness question ("how ready do you feel to lead positive change at home?") on the welcome page.',
        'Word comment 5; PPT comment (Anu)','Remove the before/after readiness check-in from the welcome (and its closing echo).',CODE,'Dev'),
  ('A3','Grammar and tense consistency across all copy.',
        'Word comment 1; Anu email','Full copy pass for consistent tense + grammar; easiest in the live/test build.',CODE,'Dev + Content'),
  ('A4','Replace "head of the family" with "his role in the family changes" (avoid reinforcing hierarchy).',
        'Word body (x2)','Change wording on the changing-role screen.',CODE,'Dev'),
 ]),
 ('B. Sound & audio', [
  ('B1','Align audio with outcomes — uplifting for positive choices, more sombre/reflective for less-healthy ones.',
        'Word body; Victor (email)','Gently differentiate outcome cues; reconcile with the deliberately minimal audio.',DECIDE,'Dev + you'),
  ('B2','Softer, subtler background music.','Word body','Already minimised the audio.',DONE,'—'),
  ('B3','Sound on/off toggle.','Word body','Already present on every screen.',DONE,'—'),
 ]),
 ('C. Reflection & participation', [
  ('C1','Answer reflective questions and download responses as a PDF — "viewed very positively."',
        'Word body; comment 3 area','Already shipped (personalised journey summary).',DONE,'—'),
  ('C2','Let players enter their own challenges / lived experiences.',
        'Word body; comment 2','Provide drop-down options in a future version.',FUTURE,'Product'),
  ('C3','End-of-game space to share experiences of racism, discrimination, settlement.',
        'Word body; comment 3','Future; raises safety/moderation questions for DV content — scope with PVAW.',FUTURE,'PVAW'),
 ]),
 ('D. Character focus & navigation', [
  ('D1','Make it clearer when the player acts AS Orion vs. as a community leader.',
        'Word body; comment 4','Resolved by A1 — Orion-only; the "community leader" second point is not required.',CODE,'Dev'),
  ('D2','Strengthen the transition from Orion’s story into the rebuild/reflection section.',
        'Word body','Add a clearer, Orion-focused bridge into the rebuild.',CODE,'Dev'),
  ('D3','Explain that some options are intentionally realistic-but-unhealthy — not "correct" answers.',
        'Word body; comment 6','Add short explainer text + note the possible impact of those choices.',CODE,'Dev'),
  ('D4','Reword "he digs in" — may not be universally understood.',
        'Word body; slide 18','Choose plainer wording for the escalation screen title/body.',CODE,'Dev'),
  ('D5','Character customisation (hairstyle, clothing, colour).',
        'Word body','Nice-to-have; larger design feature.',FUTURE,'Design'),
 ]),
 ('E. Visuals & text density', [
  ('E1','Reduce text on the opening screen (especially for mobile).',
        'Word body (x2)','Trim the welcome copy — aligns with the mobile/readability note.',CODE,'Dev'),
  ('E2','Make the black cloud more prominent; use the thought-bubble device more throughout.',
        'Word body','Illustration change in the scene engine.',DESIGN,'Design/Dev'),
  ('E3','Add icons / visual supports for lower-English users (migration, loss of status).',
        'Word body','Add supporting icons on the pressure screens.',DESIGN,'Design'),
 ]),
 ('F. Screen-specific', [
  ('F1','Money worries: show Orion actively job-seeking with repeated rejection and overseas qualifications not recognised — structural barriers, not personal failure. Keep employment visible throughout.',
        'PPT comment 102 (Ali); Word body','Copy now (structural framing); supporting visuals later.',CODE,'Dev + Design'),
  ('F2','A changing role: include Orion in the image but visually smaller than his wife and children.',
        'PPT comment 103 (Ali); Word body','Illustration change on the changing-role screen.',DESIGN,'Design'),
  ('F3','Migration: show two pressure bubbles at once (migration + unemployment); add "Orion is not sure where to go and he has a choice to make."',
        'Word body','Note tension with the earlier EAL "one pressure per screen" split — confirm direction.',DECIDE,'Dev + you'),
  ('F4','Add belief examples before the belief step; normalise help-seeking earlier (add a "seek help" option).',
        'Word body','Content addition; needs PVAW check (keep understanding ≠ excusing).',CONTENT,'Content + PVAW'),
  ('F5','"Who feels it?" → retitle "Family Experiences."',
        'Word body','Rename the impact screen.',CODE,'Dev'),
  ('F6','Add "Orion’s Reasoning" for each behaviour option (they supplied the reasoning).',
        'Word body','Content; handle carefully so it explains without excusing — PVAW check.',CONTENT,'Content + PVAW'),
  ('F7','Belief step: add more belief options and explore the drivers in greater depth (options supplied).',
        'Word body','Content expansion of the "soil"; PVAW check.',CONTENT,'Content + PVAW'),
  ('F8','Clarify whether all three response pathways lead to the same page or to different outcomes.',
        'Word body','Answer for Ali: confirm the branching behaviour.',PROCESS,'Dev → Ali'),
  ('F9','Belief step: strengthen "pressures remain the same, but we can choose different responses."',
        'Word body','Copy tweak.',CODE,'Dev'),
  ('F10','Rebuild option: add "Making decisions together while remaining responsible."',
        'Word body; slide 11','Add option; PVAW check on wording.',CONTENT,'Content + PVAW'),
  ('F11','Show the benefits of healthier choices (more confidence, less stress, stronger relationships).',
        'Word body; slide 12','Copy addition on the positive-change screen.',CODE,'Dev'),
  ('F12','Make the key message explicit: pressures may continue, but seeking support builds healthier relationships. Could this be shown via images (shared decisions, counselling)?',
        'Word body; slide 13','Copy now; supporting images later.',CODE,'Dev + Design'),
  ('F13','Outcome: include "small changes still make a meaningful difference; positive change doesn’t require major actions."',
        'Word body; slide 15','Copy addition on the healing/mixed outcome.',CODE,'Dev'),
  ('F14','"What we learn" — the reflection questions need to be changed.',
        'Word body; slide 20','Revise the four reflection questions.',CONTENT,'Content'),
  ('F15','Leadership pledge: change the name; make commitments specific, visible and time-bound (e.g. "This month, I will…").',
        'Word body; slide 22','Rename (Orion-focused) + add a time-bound prompt.',CODE,'Dev'),
 ]),
 ('G. Process & next steps', [
  ('G1','Conduct user testing with men’s groups before launch.','Word body','Plan a testing round.',PROCESS,'you / AMES'),
  ('G2','Provide drop-down options for adding future versions.','Word comment 2','Design for a later version.',FUTURE,'Product'),
  ('G3','Anu to review the wireframe / beta next; Ali to answer remaining questions.','Anu email','Deliver an updated test build for review.',PROCESS,'Dev'),
 ]),
]

# ---------------- Word ----------------
doc = Document()
sec = doc.sections[0]; sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Inches(11.69), Inches(8.27)
for m in ('left_margin','right_margin'): setattr(sec, m, Inches(0.5))
st = doc.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(9)

h = doc.add_paragraph(); r=h.add_run('Shaping Change — PAG Feedback Tracker'); r.bold=True; r.font.size=Pt(18)
sub = doc.add_paragraph(); rs=sub.add_run('Consolidated from the PAG Word document (body + 6 margin comments), the PowerPoint review comments, and Anu’s covering email. '
    'Status reflects the current build. Key decision from the comments: keep the focus on Orion for this version and drop the player-as-community-leader framing.')
rs.italic=True; rs.font.size=Pt(9.5); rs.font.color.rgb=RGBColor(0x55,0x55,0x55)
leg = doc.add_paragraph(); lr=leg.add_run('Status key:  Done · To do — code · To do — design · To do — content (PVAW) · Needs decision · Future version · Process')
lr.font.size=Pt(8.5); lr.font.color.rgb=RGBColor(0x70,0x70,0x70)

COLS=['ID','Feedback item','Source','Action / decision','Status','Owner']
WID=[Inches(0.4),Inches(3.7),Inches(1.35),Inches(3.05),Inches(1.15),Inches(0.95)]
for theme, items in THEMES:
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(2)
    tr=p.add_run(theme); tr.bold=True; tr.font.size=Pt(12); tr.font.color.rgb=RGBColor(0x1a,0x1a,0x1a)
    tbl=doc.add_table(rows=1, cols=len(COLS)); tbl.style='Table Grid'; tbl.alignment=WD_TABLE_ALIGNMENT.LEFT
    tbl.autofit=False
    for j,(c,w) in enumerate(zip(COLS,WID)):
        cell=tbl.rows[0].cells[j]; cell.width=w
        rn=cell.paragraphs[0].add_run(c); rn.bold=True; rn.font.size=Pt(8.5); rn.font.color.rgb=RGBColor(0xff,0xff,0xff)
        sh=cell._tc.get_or_add_tcPr()
        from docx.oxml.ns import qn; from docx.oxml import OxmlElement
        shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'2A1F10'); sh.append(shd)
    for (iid,fb,src,act,statv,own) in items:
        cells=tbl.add_row().cells
        vals=[iid,fb,src,act,statv,own]
        for j,(val,w) in enumerate(zip(vals,WID)):
            cells[j].width=w
            par=cells[j].paragraphs[0]; rn=par.add_run(val); rn.font.size=Pt(8.5)
            if j==0: rn.bold=True
            if j==4: rn.bold=True; rn.font.color.rgb=COLOR.get(statv,RGBColor(0,0,0))

out_docx=os.path.join(ROOT,'docs','Shaping-Change-PAG-Feedback-Tracker.docx')
doc.save(out_docx)

# ---------------- Markdown ----------------
lines=['# Shaping Change — PAG Feedback Tracker','',
 '*Consolidated from the PAG Word document (body + 6 margin comments), the PowerPoint review comments, and Anu’s covering email. '
 'Key decision from the comments: keep the focus on Orion for this version and drop the player-as-community-leader framing.*','',
 '**Status key:** Done · To do — code · To do — design · To do — content (PVAW) · Needs decision · Future version · Process','']
for theme, items in THEMES:
    lines += ['', f'## {theme}','', '| ID | Feedback item | Source | Action / decision | Status | Owner |','|---|---|---|---|---|---|']
    for (iid,fb,src,act,statv,own) in items:
        esc=lambda s: s.replace('|','\\|')
        lines.append(f'| {iid} | {esc(fb)} | {esc(src)} | {esc(act)} | {statv} | {own} |')
out_md=os.path.join(ROOT,'docs','PAG-Feedback-Tracker.md')
open(out_md,'w',encoding='utf-8').write('\n'.join(lines)+'\n')

n=sum(len(i) for _,i in THEMES)
print('saved', os.path.relpath(out_docx,ROOT))
print('saved', os.path.relpath(out_md,ROOT))
print(f'{n} items across {len(THEMES)} themes')
