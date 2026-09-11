#!/usr/bin/env python3
"""Minimal Markdown -> Word converter for this project's short docs.

Supports: # / ## / ### headings, paragraphs, `-` bullets, `1.` numbered lists,
simple pipe tables, `>` blockquotes, `---` rules, and inline **bold**, *italic*
and `code`. Written because pandoc is not reliably available in this environment
(it lives in /tmp, which gets cleared).

    pip install python-docx
    python3 scripts/md-to-docx.py <in.md> <out.docx>
"""
import os
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BROWN = RGBColor(0x5A, 0x3E, 0x1B)
GOLDINK = RGBColor(0x8A, 0x6A, 0x1E)
GREY = RGBColor(0x70, 0x6A, 0x60)
BODY = RGBColor(0x24, 0x20, 0x1A)

INLINE = re.compile(r'(\*\*.+?\*\*|\*[^*\n]+?\*|`[^`\n]+?`)')


def add_inline(p, text, size=11, color=BODY, base_italic=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            r = p.add_run(part[1:-1]); r.font.name = 'Consolas'
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            r = p.add_run(part[1:-1]); r.italic = True
        else:
            r = p.add_run(part)
        r.font.size = Pt(size)
        r.font.color.rgb = color
        if base_italic:
            r.italic = True
    return p


def shade(cell, hex_fill):
    el = OxmlElement('w:shd'); el.set(qn('w:val'), 'clear'); el.set(qn('w:fill'), hex_fill)
    cell._tc.get_or_add_tcPr().append(el)


def hrule(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    bd = OxmlElement('w:pBdr'); bot = OxmlElement('w:bottom')
    bot.set(qn('w:val'), 'single'); bot.set(qn('w:sz'), '6')
    bot.set(qn('w:space'), '1'); bot.set(qn('w:color'), 'C9C0B0')
    bd.append(bot); pPr.append(bd)
    return p


def convert(src, dst):
    lines = open(src, encoding='utf-8').read().split('\n')
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1.0)
        s.top_margin = s.bottom_margin = Inches(0.9)
    n = doc.styles['Normal'].font
    n.name = 'Calibri'; n.size = Pt(11)

    i = 0
    while i < len(lines):
        line = lines[i]
        st = line.strip()

        if not st:
            i += 1; continue

        # table: a run of lines beginning with '|'
        if st.startswith('|'):
            block = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                block.append(lines[i].strip()); i += 1
            rows = [[c.strip() for c in r.strip('|').split('|')] for r in block]
            if len(rows) >= 2 and set(rows[1][0].replace(' ', '')) <= set('-:'):
                header, body = rows[0], rows[2:]
            else:
                header, body = rows[0], rows[1:]
            t = doc.add_table(rows=1, cols=len(header)); t.style = 'Table Grid'
            for j, h in enumerate(header):
                cell = t.rows[0].cells[j]; cell.text = ''
                add_inline(cell.paragraphs[0], h, size=10.5, color=BROWN)
                for r in cell.paragraphs[0].runs: r.bold = True
                shade(cell, 'F2ECE0')
            for row in body:
                cells = t.add_row().cells
                for j, val in enumerate(row[:len(header)]):
                    cells[j].text = ''
                    add_inline(cells[j].paragraphs[0], val, size=10.5)
            doc.add_paragraph()
            continue

        # headings
        if st.startswith('### '):
            p = doc.add_paragraph(); add_inline(p, st[4:], size=12.5, color=GOLDINK)
            for r in p.runs: r.bold = True; r.font.name = 'Georgia'
            i += 1; continue
        if st.startswith('## '):
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12)
            add_inline(p, st[3:], size=15, color=BROWN)
            for r in p.runs: r.bold = True; r.font.name = 'Georgia'
            i += 1; continue
        if st.startswith('# '):
            p = doc.add_paragraph(); add_inline(p, st[2:], size=22, color=BROWN)
            for r in p.runs: r.bold = True; r.font.name = 'Georgia'
            i += 1; continue

        if set(st) <= set('-') and len(st) >= 3:
            hrule(doc); i += 1; continue

        # blockquote
        if st.startswith('>'):
            buf = []
            while i < len(lines) and lines[i].strip().startswith('>'):
                buf.append(lines[i].strip().lstrip('>').strip()); i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(8)
            add_inline(p, ' '.join(buf), size=11.5, color=BODY, base_italic=True)
            continue

        # bullets
        if st.startswith('- '):
            p = doc.add_paragraph(style='List Bullet'); add_inline(p, st[2:])
            i += 1; continue

        # numbered
        m = re.match(r'^(\d+)\.\s+(.*)$', st)
        if m:
            p = doc.add_paragraph(style='List Number'); add_inline(p, m.group(2))
            i += 1; continue

        # paragraph (join wrapped lines)
        buf = [st]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
                r'^(#{1,3}\s|-\s|\d+\.\s|\||>|---)', lines[i].strip()):
            buf.append(lines[i].strip()); i += 1
        p = doc.add_paragraph(); add_inline(p, ' '.join(buf))

    doc.save(dst)
    return dst


if __name__ == '__main__':
    if len(sys.argv) < 3:
        sys.exit('usage: md-to-docx.py <in.md> <out.docx>')
    out = convert(sys.argv[1], sys.argv[2])
    print('saved', os.path.relpath(out))
