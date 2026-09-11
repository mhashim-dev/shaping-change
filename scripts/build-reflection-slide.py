#!/usr/bin/env python3
"""Build a single review slide explaining the optional CLOSING REFLECTION (before/after).
Matches the Screen Designs deck + the landing-page review slide, so it drops straight in.

    pip install python-pptx
    python3 scripts/build-reflection-slide.py   # -> docs/Shaping-Change-Closing-Reflection-Slide.pptx
"""
import os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GOLD = RGBColor(0xF0, 0xC6, 0x6A); INK = RGBColor(0x2A, 0x1F, 0x10)
DARK = RGBColor(0x1B, 0x14, 0x0C); PANEL = RGBColor(0x17, 0x11, 0x0A)
CREAM = RGBColor(0xF5, 0xEF, 0xE2); MUTE = RGBColor(0xCF, 0xC7, 0xB8)
LINEC = RGBColor(0x4A, 0x3E, 0x2A); GREEN = RGBColor(0x9F, 0xD4, 0xA6)
SERIF, SANS = 'Georgia', 'Helvetica'
W, H = 12192000, 6858000

prs = Presentation(); prs.slide_width = Emu(W); prs.slide_height = Emu(H)
s = prs.slides.add_slide(prs.slide_layouts[6])


def tb(x, y, w, h, anchor=None):
    t = s.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h)).text_frame
    t.word_wrap = True
    if anchor is not None: t.vertical_anchor = anchor
    return t


def para(tf, runs, *, align=PP_ALIGN.LEFT, before=0, after=6, spacing=1.12, reuse=False):
    p = tf.paragraphs[0] if (reuse and len(tf.paragraphs) == 1 and not tf.paragraphs[0].runs) else tf.add_paragraph()
    p.alignment = align
    if before: p.space_before = Pt(before)
    p.space_after = Pt(after); p.line_spacing = spacing
    for (text, size, font, color, bold, italic) in runs:
        r = p.add_run(); r.text = text
        r.font.size = Pt(size); r.font.name = font; r.font.bold = bold; r.font.italic = italic
        r.font.color.rgb = color
    return p


def fill(x, y, w, h, color, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE):
    sp = s.shapes.add_shape(shape, Emu(x), Emu(y), Emu(w), Emu(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb = line; sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    return sp


def card(x, tag, q, faces, note, faces_size=30):
    CW, CY, CH = 4820000, 1720000, 3150000
    c = fill(x, CY, CW, CH, PANEL, line=GOLD, lw=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    t = tb(x + 320000, CY + 260000, CW - 640000, CH - 520000)
    para(t, [(tag, 12, SANS, GOLD, True, False)], after=10, reuse=True)
    para(t, [(q, 15, SERIF, CREAM, False, True)], after=14, spacing=1.16)
    para(t, [(faces, faces_size, SANS, CREAM, False, False)], after=12, align=PP_ALIGN.CENTER)
    para(t, [(note, 12, SANS, MUTE, False, True)], after=0, spacing=1.12)
    return x, CY, CW, CH


# background + header
fill(0, 0, W, H, DARK)
eb = tb(520000, 300000, 11000000, 460000)
para(eb, [('SHAPING CHANGE    ·    CLOSING REFLECTION    ·    FOR REVIEW', 12, SANS, GOLD, True, False)], after=0, reuse=True)
fill(520000, 740000, W - 1040000, 9000, LINEC)
ti = tb(520000, 850000, 11150000, 900000)
para(ti, [('The optional reflection — a before-and-after check-in', 24, SERIF, CREAM, True, True)], after=3, reuse=True)
para(ti, [('Optional and private. Shown at the start as a baseline, then again at the end so participants '
           'can see how their own readiness moved.', 14, SANS, MUTE, False, True)], after=0, spacing=1.14)

# two cards + arrow
ax = 680000
card(ax, '①   AT THE START', 'How ready do you feel to lead positive change at home?',
     '😟    🙁    😐    🙂    😃', 'the “before” baseline — optional (“Prefer not to say” is fine)')
bx = 6692000
card(bx, '②   AT THE END  ·  closing screen', 'You started here — and finished here:',
     '😟    →    😃', 'the shift — shown near the pathways / next steps', faces_size=34)

# arrow between the two cards
arrow = tb(ax + 4820000, 2650000, bx - (ax + 4820000), 900000, anchor=MSO_ANCHOR.MIDDLE)
para(arrow, [('→', 40, SANS, GOLD, True, False)], align=PP_ALIGN.CENTER, after=0, reuse=True)
lbl = tb(ax + 4820000, 3450000, bx - (ax + 4820000), 400000)
para(lbl, [('across the\nactivity', 10.5, SANS, MUTE, False, True)], align=PP_ALIGN.CENTER, after=0, reuse=True)

# discussion note
nb = fill(680000, 5150000, 10832000, 1150000, PANEL, line=LINEC, lw=1)
nt = tb(1000000, 5290000, 10200000, 870000)
para(nt, [('FOR DISCUSSION (with Ali)', 11, SANS, GREEN, True, False)], after=6, reuse=True)
para(nt, [('Keep the before-and-after as is — it lets participants see their start-to-finish shift; ', 14, SANS, CREAM, False, False),
          ('or', 14, SANS, GOLD, True, True),
          (' move to a single reflection at the end near the pathways (simpler, but we’d lose the start-to-finish comparison).',
           14, SANS, CREAM, False, False)], after=0, spacing=1.16)

s.notes_slide.notes_text_frame.text = (
    'The optional reflection is a before-and-after self-check-in.\n\n'
    'At the START (intro): "How ready do you feel to lead positive change at home?" — five faces, or '
    '"Prefer not to say". This is the baseline. It is optional and never blocks progress.\n\n'
    'At the END (closing / support screen, deck slide 23): the same question, shown as the shift — e.g. '
    '😟 → 😃 — so the participant sees how their own readiness moved from start to finish.\n\n'
    'It is private — nothing is collected or sent anywhere; the rating stays on the person\'s device.\n\n'
    'For discussion with Ali: keep the before/after (current), or move to a single reflection at the end '
    'near the pathways. The trade-off: a single end reflection is simpler, but there is no baseline to '
    'compare against, so we lose the "start to finish" shift.')

out = os.path.join(ROOT, 'docs', 'Shaping-Change-Closing-Reflection-Slide.pptx')
prs.core_properties.title = 'Shaping Change — Closing Reflection (for review)'
prs.core_properties.author = 'Shaping Change'
prs.save(out)
print('saved', os.path.relpath(out, ROOT), '(%d slide)' % len(prs.slides._sldIdLst))
